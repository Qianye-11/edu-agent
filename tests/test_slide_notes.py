# -*- coding: utf-8 -*-
"""备注（R16）：编辑区下方写讲稿 → 导出后 WPS 的备注栏能看到。

为什么单独一个文件：备注要跨三条导出路（新版式 / 旧版式 / 模板填充）落进 pptx，
每条路都要**读回来核对**才能在验收里说"WPS 里能看到" —— 只调赋值不读回就是自说自话。
"""

from __future__ import annotations

import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

_PROJ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(_PROJ / "src"))

from edu_agent import slide_deck                          # noqa: E402
from edu_agent import slide_fill as sf                    # noqa: E402
from edu_agent import slide_template as st                # noqa: E402
from edu_agent import wps_export                          # noqa: E402

SAMPLE_PPTX = _PROJ / "content" / "plans" / "完整课时模板（默认）.pptx"
DECK = [
    {"layout": "cover", "name": "3.3 幂函数", "cover_title": "3.3 幂函数", "notes": "开场：问大家见过哪些幂"},
    {"name": "情境导入", "minutes": "5", "bullets": ["投影五个实例"],
     "notes": "先问面积怎么算，再引出负指数"},
    {"name": "概念建构", "minutes": "12", "bullets": ["形如 y=x^α"]},
]


def _notes_of(path) -> list[str]:
    from pptx import Presentation
    out = []
    for s in Presentation(str(path)).slides:
        try:
            out.append((s.notes_slide.notes_text_frame.text or "").strip())
        except Exception:  # noqa: BLE001
            out.append("")
    return out


class NotesHelpersTest(unittest.TestCase):
    def test_slide_notes_pads_and_clamps(self):
        got = slide_deck.slide_notes(DECK, 5)
        self.assertEqual(len(got), 5, "缺的页补空串（要跟导出的页数对齐）")
        self.assertEqual(got[0], "开场：问大家见过哪些幂")
        self.assertEqual(got[2], "")
        self.assertEqual(len(slide_deck.slide_notes([{"notes": "字" * 5000}], 1)[0]), 2000,
                         "上限 2000（跟 clean_slides 一致）")

    def test_clean_slides_keeps_notes(self):
        got = slide_deck.clean_slides([dict(DECK[1], notes="  改过的讲稿  ")])[0]
        self.assertEqual(got["notes"], "改过的讲稿")

    def test_write_and_read_back(self):
        from pptx import Presentation
        tmp = tempfile.mkdtemp(prefix="edu_notes_")
        try:
            path = Path(tmp) / "x.pptx"
            prs = Presentation()
            for _ in range(3):
                prs.slides.add_slide(prs.slide_layouts[6])
            prs.save(str(path))
            res = slide_deck.write_notes(path, ["甲", "", "丙"])
            self.assertTrue(res["ok"], res)
            self.assertEqual(res["written"], 2, "空的备注不建备注页")
            self.assertEqual(slide_deck.notes_in_pptx(path), ["甲", "", "丙"])
            bad = slide_deck.write_notes(path, ["只有一条"])
            self.assertFalse(bad["ok"], "页数不一致要如实报错，不能写半截")
            self.assertIn("页数不一致", bad["error"])
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_verify_reports_missing_notes(self):
        from pptx import Presentation
        tmp = tempfile.mkdtemp(prefix="edu_notes2_")
        try:
            path = Path(tmp) / "y.pptx"
            prs = Presentation()
            prs.slides.add_slide(prs.slide_layouts[6])
            prs.save(str(path))
            svg = '<svg xmlns="http://www.w3.org/2000/svg"/>'
            no_notes = slide_deck.verify_pptx(path, [svg])
            self.assertNotIn("备注读回", [c["name"] for c in no_notes["checks"]],
                             "没给 notes 就不查这一条（老调用方不受影响）")
            miss = slide_deck.verify_pptx(path, [svg], notes=["写了却没落盘"])
            hit = [c for c in miss["checks"] if c["name"] == "备注读回"]
            self.assertTrue(hit and not hit[0]["ok"], miss["checks"])
            self.assertFalse(miss["ok"], "备注读不回来要判不合格")
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


class LegacyPptxNotesTest(unittest.TestCase):
    """旧版式（每环节一页）也要带备注：`pptx_native` 从 steps 里读 `notes`。"""

    def test_pptx_native_writes_notes(self):
        tmp = tempfile.mkdtemp(prefix="edu_notes3_")
        try:
            out = Path(tmp) / "legacy.pptx"
            wps_export.pptx_native("3.3 幂函数", "高一", [
                {"name": "导入", "minutes": 5, "bullets": ["一问"], "notes": "导入的讲稿"},
                {"name": "新课", "minutes": 10, "bullets": ["二问"], "notes": ""},
            ], out)
            notes = _notes_of(out)
            self.assertIn("导入的讲稿", "\n".join(notes), notes)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


@unittest.skipUnless(SAMPLE_PPTX.is_file(), "需要仓库样例 pptx 当模板")
class FillNotesTest(unittest.TestCase):
    """模板填充也带备注：老师写在幻灯片页签上的讲稿，填出来的 pptx 里读得回。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="edu_fillnotes_")
        self._old = os.environ.get("EDU_PPT_TEMPLATES")
        os.environ["EDU_PPT_TEMPLATES"] = self.tmp.name
        builtin = Path(self.tmp.name) / "_builtin"
        builtin.mkdir(parents=True)
        for f in (_PROJ / "content" / "ppt_templates" / "_builtin").glob("*.json"):
            shutil.copy(f, builtin / f.name)
        self.out = Path(self.tmp.name) / "out"
        self.tid = st.import_pptx(SAMPLE_PPTX, name="学校模板")["id"]
        self.addCleanup(self._restore)

    def _restore(self):
        if self._old is None:
            os.environ.pop("EDU_PPT_TEMPLATES", None)
        else:
            os.environ["EDU_PPT_TEMPLATES"] = self._old
        self.tmp.cleanup()

    def test_plan_carries_note_per_page(self):
        plan = sf.plan(self.tid, DECK, "备注")
        self.assertEqual(plan["pages"][0]["note"], "开场：问大家见过哪些幂")
        self.assertEqual(plan["pages"][1]["note"], "先问面积怎么算，再引出负指数")
        self.assertTrue(plan["pages"][2]["note"], "没写讲稿的页也要有时长备注（信息不丢）")
        self.assertIn("12", plan["pages"][2]["note"])

    def test_fill_writes_notes_and_reports_check(self):
        got = sf.fill_pptx(self.tid, DECK, "备注", out_dir=self.out)
        self.assertTrue(got["ok"], got.get("error"))
        notes = _notes_of(got["file"])
        self.assertIn("先问面积怎么算，再引出负指数", "\n".join(notes), notes)
        hit = [c for c in got["checks"] if c["name"] == "备注读回"]
        self.assertTrue(hit and hit[0]["ok"], got["checks"])
        self.assertFalse(hit[0]["blocking"], "备注只提示，不阻断填充")
        self.assertGreaterEqual(got["notes_written"], 2, got["notes_written"])


if __name__ == "__main__":
    unittest.main()
