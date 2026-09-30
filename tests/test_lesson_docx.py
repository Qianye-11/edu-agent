from __future__ import annotations

from docx import Document

from edu_agent.lesson_docx import export_html


def test_export_html_keeps_title_and_table(tmp_path):
    path = export_html(
        "<h1>函数单调性教案</h1><table><tbody><tr><td>教学目标</td>"
        "<td><p><b>知识与技能：</b>理解单调性定义</p></td></tr></tbody></table>",
        "函数单调性教案",
        tmp_path,
    )

    document = Document(path)

    assert path.suffix == ".docx"
    assert document.paragraphs[0].text == "函数单调性教案"
    assert document.tables[0].cell(0, 0).text == "教学目标"
    assert "理解单调性定义" in document.tables[0].cell(0, 1).text


def test_export_html_preserves_a4_lesson_plan_layout(tmp_path):
    path = export_html(
        "<h1>规范教案</h1><table><tbody>"
        "<tr><td class='lp-label'>授课教师</td><td></td><td class='lp-label'>授课班级</td>"
        "<td></td><td class='lp-label'>授课时间</td><td></td></tr>"
        "<tr><td class='lp-label lp-process-label'>教<br>学<br>过<br>程</td>"
        "<td colspan='5'><p>教师活动与学生任务</p></td></tr>"
        "</tbody></table>",
        "规范教案",
        tmp_path,
    )

    document = Document(path)
    section = document.sections[0]
    table = document.tables[0]

    assert round(section.page_width.cm, 1) == 21.0
    assert round(section.page_height.cm, 1) == 29.7
    assert table.autofit is False
    assert len(table.columns) == 6
    assert round(table.rows[1].height.cm, 1) == 11.4
    assert table.cell(1, 0).text == "教\n学\n过\n程"


def test_export_html_converts_editor_page_break_to_word_page_break(tmp_path):
    path = export_html(
        "<h1>多页教案</h1><p>第一页内容</p>"
        "<div class='lesson-page-break' contenteditable='false'><span>分页符 · 下一页</span></div>"
        "<h2>第二页内容</h2><p>继续教学活动。</p>",
        "多页教案",
        tmp_path,
    )

    document = Document(path)
    page_breaks = document.element.body.xpath(".//w:br[@w:type='page']")
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert len(page_breaks) == 1
    assert "第一页内容" in text
    assert "第二页内容" in text
    assert "分页符" not in text


def test_export_html_keeps_canvas_data_url_image(tmp_path):
    # 1×1 PNG：覆盖编辑器粘贴图片的 data URL 导出路径。
    image = (
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADElEQVR42mNk+M/wHwAF/gL+"
        "5sLhAAAAAElFTkSuQmCC"
    )
    path = export_html(
        f"<h1>含图片教案</h1><p><img src='data:image/png;base64,{image}'></p>",
        "含图片教案",
        tmp_path,
    )

    document = Document(path)

    assert len(document.inline_shapes) == 1
