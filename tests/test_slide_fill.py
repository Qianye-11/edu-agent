# -*- coding: utf-8 -*-
"""填充器（M4b / docs/19 §10）单测：映射规则、页数策略、两道校验、红线。

映射规则是**确定性**的，所以逐条对：封面→封面槽、`name`→标题槽、`bullets`→正文槽、`minutes`→页脚槽；
模板没有的角色**不硬造**；页数不匹配按 §10.2 的策略走（多→复制模板自己的内容页；少→默认删多余页）。
"""

from __future__ import annotations

import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

_PROJ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(_PROJ / "src"))

from edu_agent import slide_deck                 # noqa: E402
from edu_agent import slide_fill as sf           # noqa: E402
from edu_agent import slide_template as st       # noqa: E402

SAMPLE_PPTX = _PROJ / "content" / "plans" / "完整课时模板（默认）.pptx"
DECK = [
    {"layout": "cover", "name": "3.3 幂函数", "cover_title": "3.3 幂函数",
     "subtitle": "数学 · 高一", "meta": "1 课时"},
    {"name": "情境导入", "minutes": "5", "bullets": ["投影五个实例", "追问：共同点"]},
    {"name": "概念建构", "minutes": "12", "bullets": ["形如 y=x^α"]},
]


class _FillCase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="edu_fill_")
        self._old = os.environ.get("EDU_PPT_TEMPLATES")
        os.environ["EDU_PPT_TEMPLATES"] = self.tmp.name
        builtin = Path(self.tmp.name) / "_builtin"
        builtin.mkdir(parents=True)
        for f in (_PROJ / "content" / "ppt_templates" / "_builtin").glob("*.json"):
            shutil.copy(f, builtin / f.name)
        self.out = Path(self.tmp.name) / "out"
        self.tid = ""
        if SAMPLE_PPTX.is_file():
            self.tid = st.import_pptx(SAMPLE_PPTX, name="学校模板")["id"]
        self.addCleanup(self._restore)

    def _restore(self):
        if self._old is None:
            os.environ.pop("EDU_PPT_TEMPLATES", None)
        else:
            os.environ["EDU_PPT_TEMPLATES"] = self._old
        self.tmp.cleanup()


@unittest.skipUnless(SAMPLE_PPTX.is_file(), "需要仓库样例 pptx 当模板")
class PlanTest(_FillCase):
    def test_cover_maps_to_cover_page(self):
        p = sf.plan(self.tid, DECK, "3.3 幂函数")
        self.assertTrue(p["pages"])
        self.assertEqual(p["pages"][0]["fills"][0]["role"], "title")
        self.assertIn("3.3 幂函数", p["pages"][0]["fills"][0]["text"])

    def test_page_names_and_bullets_go_in_order(self):
        p = sf.plan(self.tid, DECK, "3.3 幂函数")
        bodies = ["".join(f["text"] for f in page["fills"]) for page in p["pages"][1:]]
        self.assertTrue(any("情境导入" in b for b in bodies))
        self.assertTrue(any("投影五个实例" in b for b in bodies))
        self.assertTrue(any("形如 y=x^α" in b for b in bodies))

    def test_fewer_pages_deletes_extra(self):
        p = sf.plan(self.tid, DECK, "3.3 幂函数")
        self.assertGreater(p["policy"].get("deleted", 0), 0)
        self.assertTrue(any("删掉" in w for w in p["warnings"]))

    def test_more_pages_copies_template_content_page(self):
        deck = DECK + [{"name": f"补充 {i}", "bullets": [f"第 {i} 条"]} for i in range(20)]
        p = sf.plan(self.tid, deck, "3.3 幂函数")
        self.assertGreater(p["policy"].get("copied", 0), 0)
        self.assertTrue(any("复制" in w for w in p["warnings"]))
        self.assertGreaterEqual(len(p["pages"]), len(deck), "页数 ≥ 会话页数（超出容量时还会再分页）")

    def test_no_slots_degrades(self):
        (Path(self.tmp.name) / "empty-tpl").mkdir()
        (Path(self.tmp.name) / "empty-tpl" / "theme.json").write_text(
            '{"name":"空模板","tokens":{}}', encoding="utf-8")
        p = sf.plan("empty-tpl", DECK, "x")
        self.assertEqual(p["pages"], [])
        self.assertTrue(any("只用配色" in d or "槽位" in d for d in p["degraded"]))


@unittest.skipUnless(SAMPLE_PPTX.is_file(), "需要仓库样例 pptx 当模板")
class FillTest(_FillCase):
    def test_fill_produces_readable_pptx_with_content(self):
        out = sf.fill_pptx(self.tid, DECK, "3.3 幂函数", out_dir=self.out)
        self.assertTrue(out["ok"], out.get("error"))
        target = Path(out["file"])
        self.assertTrue(target.is_file())
        from pptx import Presentation

        prs = Presentation(str(target))
        blob = "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes
                         if getattr(sh, "has_text_frame", False))
        for want in ("3.3 幂函数", "情境导入", "投影五个实例", "形如 y=x^α"):
            with self.subTest(text=want):
                self.assertIn(want, blob)

    def test_two_checks_present_and_blocking_flag(self):
        out = sf.fill_pptx(self.tid, DECK, "3.3 幂函数", out_dir=self.out)
        names = {c["name"]: c for c in out["checks"]}
        self.assertIn("内容零丢失", names)
        self.assertIn("容量自查", names)
        self.assertTrue(names["内容零丢失"]["ok"])
        self.assertFalse(names["容量自查"]["blocking"], "容量自查只提示，不阻断")

    def test_template_original_untouched(self):
        tpl = st.root() / self.tid / "template.pptx"
        before = (tpl.stat().st_size, tpl.stat().st_mtime)
        sf.fill_pptx(self.tid, DECK, "3.3 幂函数", out_dir=self.out)
        after = (tpl.stat().st_size, tpl.stat().st_mtime)
        self.assertEqual(before, after, "模板原稿只读：填充在副本上做")

    def test_never_writes_content_plans(self):
        plans = _PROJ / "content" / "plans"
        before = {p.name for p in plans.glob("*")} if plans.is_dir() else set()
        sf.fill_pptx(self.tid, DECK, "3.3 幂函数", out_dir=self.out)
        after = {p.name for p in plans.glob("*")} if plans.is_dir() else set()
        self.assertEqual(before, after)

    def test_long_text_is_flagged_but_not_cut(self):
        """R6 新策略：**不允许"容量不足"** —— 由"适度压缩 + 另起一页"消化，且内容不丢。

        2026-09-27 调整：容量改成按"允许缩到的最小字号"算之后，框能装的字**多了**，
        所以这里按"最大框容量"现算超长文本（老写法写死 190 字，已经不超容量了 → 断言失去意义）。
        """
        probe = sf.plan(self.tid, DECK, "3.3 幂函数")
        caps = [int(f.get("capacity") or 0) for page in probe["pages"] for f in page["fills"]]
        cap = max(caps) if caps else 0
        per = max(40, cap * 3 // 2)                          # 每条都比最大的框还长 50%
        bullets = ["要点" * (per // 2) for _ in range(2)]
        deck = [DECK[0], {"name": "长内容", "bullets": bullets}]
        self.assertGreater(len(bullets[0]), cap)
        p = sf.plan(self.tid, deck, "3.3 幂函数")
        self.assertFalse(any("容量" in w for w in p.get("warnings") or []), p.get("warnings"))
        self.assertTrue(p.get("adjustments"), "压缩/分页要留痕（adjustments）")
        self.assertTrue({"trim", "new-page", "too-long"} & {a.get("action") for a in p["adjustments"]},
                        p["adjustments"])
        out = sf.fill_pptx(self.tid, deck, "3.3 幂函数", out_dir=self.out)
        self.assertTrue(out["ok"], out.get("error"))
        from pptx import Presentation

        prs = Presentation(out["file"])
        blob = "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes
                         if getattr(sh, "has_text_frame", False))
        for i, b in enumerate(bullets):
            with self.subTest(bullet=i):
                self.assertIn(b[:12], blob, "每条的开头都要在（压缩只去尾巴，不整条丢）")
        self.assertTrue({c["name"]: c["ok"] for c in out["checks"]}["内容零丢失"])

    def test_unknown_template_is_explicit(self):
        out = sf.fill_pptx("没有这个模板", DECK, "x", out_dir=self.out)
        self.assertFalse(out["ok"])
        self.assertIn("没有原稿文件", out["error"])


@unittest.skipUnless(SAMPLE_PPTX.is_file(), "需要仓库样例 pptx 当模板")
class RenderWithTemplateTest(_FillCase):
    def test_imported_template_can_render_deck(self):
        tokens = st.theme_for(self.tid)
        out = slide_deck.build_deck("", slide_deck.SAMPLE_DECK, title="3.3 幂函数", tokens=tokens)
        self.assertEqual(len(out["pages"]), len(slide_deck.SAMPLE_DECK))
        self.assertIn(tokens["canvas"]["bg"], out["pages"][0])


class MultiRunRegressionTest(_FillCase):
    """回归：**一段里有多个 run** 的模板（老师从 PowerPoint 里粘来的标题常见）。

    2026-09-25 线上报错：`AttributeError: '_Run' object has no attribute '_element'`
    —— python-pptx 的元素访问器不一样（`_Run._r` / `_Paragraph._p` / 形状 `._element`），
    而仓库样例的段落只有一个 run，所以当时的单测没覆盖到这条路径。
    """

    def _multi_run_template(self) -> Path:
        """一个"像真设计稿"的模板：标题占位符里一段多 run 文字 + 第二行；
        一个**设计标签**（短文本 + 小框，如「快乐成长」）；一个**内容槽**（示例正文）。"""
        from pptx import Presentation
        from pptx.util import Inches, Pt

        src = Path(self.tmp.name) / "multi-run.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])          # Title Only：有标题占位符
        tf = slide.shapes.title.text_frame
        r1 = tf.paragraphs[0].add_run()
        r1.text = "点击此处添加标题"
        r1.font.size = Pt(32)
        r1.font.bold = True
        r2 = tf.paragraphs[0].add_run()                             # ← 第二个 run：触发删多余 run 的分支
        r2.text = "（模板原名）"
        para2 = tf.add_paragraph()
        para2.add_run().text = "（占位第二行）"                      # ← 第二段：触发删多余段落的分支
        label = slide.shapes.add_textbox(Inches(0.4), Inches(0.3), Inches(0.9), Inches(0.25))
        label.text_frame.text = "快乐成长"                            # 短文本 + 小框 → 设计标签
        body = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(0.6))
        body.text_frame.text = "模板自己的装饰文字"                    # 较长文本 → 内容槽（示例正文）
        prs.save(str(src))
        return src

    def test_multi_run_paragraph_is_handled(self):
        """回归 2026-09-25：`AttributeError: '_Run' object has no attribute '_element'`
        （python-pptx 的 `_Run` 是 `._r`；仓库样例的段落只有一个 run，所以当时没覆盖到）。"""
        tid = st.import_pptx(self._multi_run_template(), name="多run模板")["id"]
        deck = [{"layout": "cover", "name": "3.3 幂函数", "cover_title": "3.3 幂函数"}]
        out = sf.fill_pptx(tid, deck, "3.3 幂函数", out_dir=self.out)
        self.assertTrue(out["ok"], out.get("error"))
        from pptx import Presentation

        prs = Presentation(out["file"])
        blob = "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes
                         if getattr(sh, "has_text_frame", False))
        self.assertIn("3.3 幂函数", blob)
        self.assertNotIn("（模板原名）", blob, "多余 run 的文字要被换掉")
        self.assertNotIn("（占位第二行）", blob, "多余段落的文字要被换掉")

    def test_run_formatting_of_first_run_kept(self):
        """替换文字时保留第一个 run 的格式（字号/加粗），否则标题会变小。"""
        tid = st.import_pptx(self._multi_run_template(), name="多run模板")["id"]
        deck = [{"layout": "cover", "name": "封面标题", "cover_title": "封面标题"}]
        out = sf.fill_pptx(tid, deck, "封面标题", out_dir=self.out)
        self.assertTrue(out["ok"], out.get("error"))
        from pptx import Presentation

        prs = Presentation(out["file"])
        sizes = [run.font.size.pt for s in prs.slides for sh in s.shapes
                 if getattr(sh, "has_text_frame", False)
                 for p in sh.text_frame.paragraphs for run in p.runs
                 if run.text.strip() and run.font.size]
        self.assertIn(32.0, sizes, "第一个 run 的 32pt 应该留着")

    def test_unfilled_content_box_is_cleared_and_labels_kept(self):
        """策略（2026-09-27 改）：**内容槽**没填就清空（免得模板示例文字和我们的内容混在一起＝乱填）；
        **设计标签**（短文本 + 小面积，如「快乐成长」「PART 01」）一个字都不动。"""
        tid = st.import_pptx(self._multi_run_template(), name="多run模板")["id"]
        out = sf.fill_pptx(tid, [{"layout": "cover", "name": "只有标题", "cover_title": "只有标题"}],
                           "只有标题", out_dir=self.out)
        self.assertTrue(out["ok"], out.get("error"))
        from pptx import Presentation

        prs = Presentation(out["file"])
        blob = "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes
                         if getattr(sh, "has_text_frame", False))
        self.assertIn("只有标题", blob, "我们的标题要填进去")
        self.assertIn("快乐成长", blob, "设计标签不能被替换/清空")
        self.assertNotIn("模板自己的装饰文字", blob, "没填的内容槽要清掉（否则和我们的内容混在一起）")
        plan = sf.plan(tid, [{"layout": "cover", "name": "x", "cover_title": "x"}], "x")
        self.assertTrue(plan["pages"][0]["kept"], "计划里要记下保留了哪些模板文字")


class ContinuationPageTest(_FillCase):
    """续页必须**整页克隆**（R7，老师报的"填充后有空白页、模板背景都没有"）。

    老实现只深拷贝 `has_text_frame` 的形状 → 续页只剩 3 个文本框、0 张图，
    底图/装饰全丢，成品里那几页看着就是白页。
    """

    def _page_with_background(self) -> Path:
        from pptx import Presentation
        from pptx.util import Emu, Pt

        def px(v: float) -> Emu:
            return Emu(int(v * 9525))

        src = Path(self.tmp.name) / "bg-tpl.pptx"
        prs = Presentation()
        for i in range(2):                                   # 封面 + 一页内容页
            s = prs.slides.add_slide(prs.slide_layouts[6])
            pic = s.shapes.add_picture(str(self._png()), px(0), px(0), px(1280), px(720))  # 整页底图
            pic.name = "底图"
            t = s.shapes.add_textbox(px(40), px(40), px(600), px(80))
            t.text_frame.text = "模板标题占位"
            t.text_frame.paragraphs[0].runs[0].font.size = Pt(36)
            if i:                                            # 内容页多一个正文框
                b = s.shapes.add_textbox(px(40), px(200), px(900), px(300))
                b.text_frame.text = "模板正文示例"
                b.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
        prs.save(str(src))
        return src

    def _png(self) -> Path:
        from PIL import Image

        p = Path(self.tmp.name) / "bg.png"
        if not p.is_file():
            Image.new("RGB", (64, 36), (200, 120, 90)).save(p)
        return p

    def test_continuation_page_keeps_pictures_and_shapes(self):
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        tid = st.import_pptx(self._page_with_background(), name="带底图模板")["id"]
        deck = [{"layout": "cover", "name": "封面", "cover_title": "封面"}] + \
               [{"name": f"第{i}页", "bullets": [f"内容{i}"]} for i in range(2, 6)]
        out = sf.fill_pptx(tid, deck, "续页检查", out_dir=self.out)
        self.assertTrue(out["ok"], out.get("error"))
        prs = Presentation(out["file"])
        src = Presentation(st.root() / tid / "template.pptx")
        src_slides = list(src.slides)
        self.assertGreater(len(prs.slides._sldIdLst), len(src_slides), "这条用例要走到「复制页」分支")
        for i, page in enumerate(out["plan"]["pages"], 1):
            donor = src_slides[page["template_index"]]
            want_pics = sum(1 for s in donor.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
            with self.subTest(page=i, donor=page["template_index"] + 1):
                self.assertGreater(want_pics, 0, "母版必须有图，否则这条用例没意义")
                slide = list(prs.slides)[i - 1]
                self.assertEqual(len(list(slide.shapes)), len(list(donor.shapes)),
                                 "每页都要和它自己的母版一样多的形状（底图/装饰不能丢）")
                pics = sum(1 for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
                self.assertEqual(pics, want_pics, "底图要在")
        names = {c["name"]: c["ok"] for c in out["checks"]}
        self.assertTrue(names["续页设计没丢"])
        self.assertTrue(names["没有空白页"])

    def test_clone_is_taken_before_writing(self):
        """克隆母版必须是"还没被填过"的原页，否则续页会带着上一页的文字。"""
        from pptx import Presentation

        tid = st.import_pptx(self._page_with_background(), name="带底图模板")["id"]
        deck = [{"layout": "cover", "name": "封面", "cover_title": "封面"}] + \
               [{"name": f"第{i}页", "bullets": [f"内容{i}"]} for i in range(2, 6)]
        out = sf.fill_pptx(tid, deck, "克隆顺序", out_dir=self.out)
        prs = Presentation(out["file"])
        blob = []
        for slide in prs.slides:
            blob.append("\n".join(sh.text_frame.text for sh in slide.shapes
                                  if getattr(sh, "has_text_frame", False)))
        for i, text in enumerate(blob, 1):
            with self.subTest(page=i):
                self.assertNotIn("模板标题占位", text, "母版的示例文字不该留在成品里")
                self.assertNotIn("模板正文示例", text)


class PaginationAndTrimTest(_FillCase):
    """分页与压缩（R7）：**按每页自己的槽**分页，且不把正文压成几个字。"""

    def _first_page_small_second_big_template(self) -> Path:
        """第 1 张内容页只有 1 个 20 字的窄框，第 2 张有大框 —— 老实现拿第 1 张的容量分所有页。"""
        from pptx import Presentation
        from pptx.util import Emu, Pt

        def px(v: float) -> Emu:
            return Emu(int(v * 9525))

        src = Path(self.tmp.name) / "caps-tpl.pptx"
        prs = Presentation()
        cover = prs.slides.add_slide(prs.slide_layouts[6])
        cover.shapes.add_textbox(px(40), px(40), px(600), px(80)).text_frame.text = "封面标题占位"
        small = prs.slides.add_slide(prs.slide_layouts[6])       # 窄框页（像目录页）
        st1 = small.shapes.add_textbox(px(40), px(40), px(300), px(60))
        st1.text_frame.text = "模板标题"
        st1.text_frame.paragraphs[0].runs[0].font.size = Pt(30)
        sb = small.shapes.add_textbox(px(40), px(200), px(200), px(60))
        sb.text_frame.text = "示例"
        sb.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
        big = prs.slides.add_slide(prs.slide_layouts[6])         # 大框页
        bt = big.shapes.add_textbox(px(40), px(40), px(600), px(80))
        bt.text_frame.text = "模板标题"
        bt.text_frame.paragraphs[0].runs[0].font.size = Pt(36)
        bb = big.shapes.add_textbox(px(40), px(200), px(1100), px(400))
        bb.text_frame.text = "示例正文"
        bb.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
        prs.save(str(src))
        return src

    def test_long_bullet_is_not_squeezed_into_a_narrow_box(self):
        from pptx import Presentation

        tid = st.import_pptx(self._first_page_small_second_big_template(), name="容量模板")["id"]
        long_text = "这是一条比较长的正文内容，用来验证不会被压成几个字塞进窄框里。" * 2
        deck = [{"layout": "cover", "name": "封面", "cover_title": "封面"},
                {"name": "窄框页", "bullets": ["短句"]},
                {"name": "长文页", "bullets": [long_text]}]
        out = sf.fill_pptx(tid, deck, "不压字数", out_dir=self.out)
        self.assertTrue(out["ok"], out.get("error"))
        prs = Presentation(out["file"])
        texts = [sh.text_frame.text for s in prs.slides for sh in s.shapes
                 if getattr(sh, "has_text_frame", False)]
        hits = [t for t in texts if long_text[:20] in t]
        self.assertTrue(hits, "长文要原样落进某个框（不被砍成短句）")
        self.assertGreaterEqual(len(hits[0]), len(long_text) * sf.LOSS_LIMIT,
                                "压缩不能砍掉一半以上（老师要求：不要容量不足）")

    def test_pagination_uses_each_page_own_capacity(self):
        tid = st.import_pptx(self._first_page_small_second_big_template(), name="容量模板")["id"]
        deck = [{"layout": "cover", "name": "封面", "cover_title": "封面"},
                {"name": "第一页", "bullets": ["甲乙丙丁戊己庚辛壬癸"]},
                {"name": "第二页", "bullets": ["子丑寅卯辰巳午未申酉"]}]
        p = sf.plan(tid, deck, "分页")
        pages = [pg for pg in p["pages"]]
        self.assertEqual(len(pages), 3, "三张会话页对应三张模板页，不该因为按错容量分页而变多")
        self.assertFalse([a for a in p.get("adjustments") or [] if a.get("action") == "new-page"],
                         "短句在规定容量内，不该另起页")


class DesignBakedFillTest(_FillCase):
    """设计稿模板（无占位符）的填充效果（R7）：正文进大框、图形标签不动。"""

    def _card_template(self) -> Path:
        from pptx import Presentation
        from pptx.util import Emu, Pt

        def px(v: float) -> Emu:
            return Emu(int(v * 9525))

        src = Path(self.tmp.name) / "cards.pptx"
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[6])
        t = s.shapes.add_textbox(px(440), px(7), px(400), px(100))
        t.text_frame.text = "图形的奥秘"
        t.text_frame.paragraphs[0].runs[0].font.size = Pt(50)
        for i in range(4):
            x = 29 + i * 319
            h = s.shapes.add_textbox(px(x), px(331), px(264), px(36))
            h.text_frame.text = f"卡片标题{i + 1}"
            h.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
            b = s.shapes.add_textbox(px(x), px(380), px(264), px(160))
            b.text_frame.text = "卡片正文示例文字" * 6
            b.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
        lab = s.shapes.add_textbox(px(87), px(139), px(58), px(68))
        lab.text_frame.text = "圆"
        lab.text_frame.paragraphs[0].runs[0].font.size = Pt(36)
        prs.save(str(src))
        return src

    def test_body_goes_into_the_roomy_box_and_labels_are_kept(self):
        from pptx import Presentation

        tid = st.import_pptx(self._card_template(), name="卡片模板")["id"]
        bullets = ["这是一整段正文，应该落在容量大的正文框里，而不是被塞进小标题框再压成几个字。",
                   "第二条正文也要完整。"]
        deck = [{"layout": "cover", "name": "封面", "cover_title": "封面"},
                {"name": "卡片页", "bullets": bullets}]
        out = sf.fill_pptx(tid, deck, "卡片填充", out_dir=self.out)
        self.assertTrue(out["ok"], out.get("error"))
        prs = Presentation(out["file"])
        texts = [sh.text_frame.text for s in prs.slides for sh in s.shapes
                 if getattr(sh, "has_text_frame", False)]
        self.assertTrue(any(bullets[0][:18] in t for t in texts), "长正文要完整落进大框")
        self.assertTrue(any(t.strip() == "圆" for t in texts), "图形标签一个字都不能动")
        self.assertFalse(any("卡片标题" in t for t in texts), "没填的内容槽要清掉（不留模板示例文字）")


class NoCoverDeckMappingTest(_FillCase):
    """回归 2026-09-27：模板有封面页、会话没有封面时，**不能把正文写进封面页**。

    体检脚本发现：老实现"计划第 i 页就写第 i 张"，于是正文落到模板封面上，
    而且没用上的封面/结尾页没被删掉（`页数符合策略` 判为 False）。现在按 `template_index` 落位 + 重排。
    """

    def _two_layout_template(self) -> Path:
        from pptx import Presentation

        src = Path(self.tmp.name) / "cover-and-content.pptx"
        prs = Presentation()
        s0 = prs.slides.add_slide(prs.slide_layouts[0])          # Title Slide（封面）
        s0.shapes.title.text_frame.text = "封面占位"
        for _ in range(2):
            s = prs.slides.add_slide(prs.slide_layouts[1])        # Title and Content
            s.shapes.title.text_frame.text = "标题占位"
            s.placeholders[1].text_frame.text = "正文占位"
        prs.save(str(src))
        return src

    def test_deck_without_cover_skips_template_cover(self):
        tid = st.import_pptx(self._two_layout_template(), name="封面+内容模板")["id"]
        deck = [{"name": "情境导入", "minutes": "5", "bullets": ["投影五个实例"]},
                {"name": "小结", "minutes": "3", "bullets": ["三条结论"]}]
        p = sf.plan(tid, deck, "3.3 幂函数")
        pages, _ = sf.template_pages(tid)
        self.assertEqual(p["policy"].get("deleted"), len(pages) - len(p["pages"]),
                         "没用到的模板页（含封面）都要计入删除")
        out = sf.fill_pptx(tid, deck, "3.3 幂函数", out_dir=self.out)
        self.assertTrue(out["ok"], out.get("error"))
        from pptx import Presentation

        prs = Presentation(out["file"])
        self.assertEqual(len(prs.slides._sldIdLst), len(p["pages"]), "产物页数 = 计划页数")   # noqa: SLF001
        blob = "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes
                         if getattr(sh, "has_text_frame", False))
        self.assertIn("情境导入", blob)
        self.assertIn("投影五个实例", blob)
        self.assertIn("小结", blob)
        self.assertNotIn("封面占位", blob, "模板封面页应该被删掉，而不是留着当正文页")

    def test_deck_with_cover_uses_template_cover(self):
        tid = st.import_pptx(self._two_layout_template(), name="封面+内容模板")["id"]
        deck = [{"layout": "cover", "name": "3.3 幂函数", "cover_title": "3.3 幂函数",
                 "subtitle": "数学 · 高一"},
                {"name": "情境导入", "bullets": ["投影五个实例"]}]
        p = sf.plan(tid, deck, "3.3 幂函数")
        out = sf.fill_pptx(tid, deck, "3.3 幂函数", out_dir=self.out)
        self.assertTrue(out["ok"], out.get("error"))
        from pptx import Presentation

        prs = Presentation(out["file"])
        self.assertEqual(len(prs.slides._sldIdLst), len(p["pages"]))         # noqa: SLF001
        blob = "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes
                         if getattr(sh, "has_text_frame", False))
        self.assertIn("3.3 幂函数", blob)
        self.assertIn("情境导入", blob)

    def test_more_slides_than_template_pages(self):
        """会话比模板页多：复制模板的内容页（母版继承不变），页数仍对得上。"""
        tid = st.import_pptx(self._two_layout_template(), name="封面+内容模板")["id"]
        deck = [{"name": f"环节 {i}", "bullets": [f"第 {i} 条"]} for i in range(6)]
        p = sf.plan(tid, deck, "3.3 幂函数")
        self.assertGreater(p["policy"].get("copied", 0), 0)
        out = sf.fill_pptx(tid, deck, "3.3 幂函数", out_dir=self.out)
        self.assertTrue(out["ok"], out.get("error"))
        from pptx import Presentation

        prs = Presentation(out["file"])
        checks = {c["name"]: c["ok"] for c in out["checks"]}
        self.assertTrue(checks["页数符合策略"], out["checks"])
        self.assertEqual(len(prs.slides._sldIdLst), len(deck))               # noqa: SLF001
        blob = "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes
                         if getattr(sh, "has_text_frame", False))
        for i in range(6):
            with self.subTest(page=i):
                self.assertIn(f"环节 {i}", blob)


class AgentApiTest(_FillCase):
    """R6-3：给外部 agent 的接口（不做 skill）—— 读端 plan、写端 mapping / 一步出稿。"""

    def setUp(self):
        super().setUp()
        import tempfile as _tf

        self._old_data = os.environ.get("EDU_DATA_DIR")
        os.environ["EDU_DATA_DIR"] = _tf.mkdtemp(prefix="edu_agent_api_")
        from edu_agent import capabilities as cap
        from edu_agent import web_server as ws

        self.ws, self.cap = ws, cap
        cap._CACHE.clear()
        self.addCleanup(self._restore_data)

    def _restore_data(self):
        if self._old_data is None:
            os.environ.pop("EDU_DATA_DIR", None)
        else:
            os.environ["EDU_DATA_DIR"] = self._old_data
        self.cap._CACHE.clear()

    @unittest.skipUnless(SAMPLE_PPTX.is_file(), "需要仓库样例 pptx 当模板")
    def test_plan_payload_schema_and_stability(self):
        a = sf.plan_payload(self.tid, DECK, "3.3 幂函数")
        b = sf.plan_payload(self.tid, DECK, "3.3 幂函数")
        self.assertTrue(a["ok"], a.get("error"))
        self.assertEqual(a["schema"], sf.SCHEMA)
        self.assertEqual(a, b, "同输入同输出（agent 要能据此做判断）")
        self.assertIn("slots", a["template"])
        roles = {s["role"] for p in a["template"]["slots"] for s in p["slots"]}
        self.assertTrue(roles <= {"title", "body", "decoration"}, roles)

    @unittest.skipUnless(SAMPLE_PPTX.is_file(), "需要仓库样例 pptx 当模板")
    def test_apply_mapping_sanitizes_and_trims(self):
        p = sf.plan(self.tid, DECK, "3.3 幂函数")
        first = p["pages"][0]
        ti = first["template_index"]
        cap = int((first["fills"][0].get("capacity") or 0))
        long_text = "字" * (cap + 20) if cap else "字" * 50
        out = sf.apply_mapping(p, [
            {"template_index": ti, "fills": [{"role": first["fills"][0]["role"], "text": long_text}]},
            {"template_index": 9999, "fills": [{"role": "title", "text": "不该生效"}]},      # 不存在的页
            {"template_index": ti, "fills": [{"role": "不存在的角色", "text": "不该生效"}]},   # 不存在的角色
        ])
        self.assertEqual(out["mapping_applied"], 1, "只认已有的页与角色")
        got = out["pages"][0]["fills"][0]["text"]
        self.assertLessEqual(len(got), cap or len(got), "超容量要被压到放得下")
        self.assertTrue(any(a.get("by") == "agent-mapping" for a in out.get("adjustments") or []))
        self.assertNotIn("不该生效", str(out))

    @unittest.skipUnless(SAMPLE_PPTX.is_file(), "需要仓库样例 pptx 当模板")
    def test_deck_one_shot_endpoint(self):
        import unittest.mock as mock

        called: dict = {}

        def fake_fill(tid, slides, title="", **kw):
            called.update({"tid": tid, "slides": slides, "title": title, "out": kw.get("out_dir")})
            return {"ok": True, "filename": "x.pptx", "file": "d/x.pptx", "outDir": "d",
                    "download_url": "/d/x", "checks": [], "pages": len(slides)}

        with mock.patch("edu_agent.slide_fill.fill_pptx", side_effect=fake_fill):
            r = self.ws.api_slide_deck(self.ws.SlideDeckReq(
                lesson={"title": "3.3 幂函数", "slides": DECK}, template_id=self.tid,
                outDir=str(self.out), confirm=True))
        self.assertTrue(r["ok"], r.get("error"))
        self.assertEqual(called["tid"], self.tid)
        self.assertEqual(r["schema"], sf.SCHEMA)

    def test_deck_without_template_uses_theme_path(self):
        import unittest.mock as mock

        with mock.patch.object(self.ws.slide_deck, "export_deck_pptx",
                               return_value={"ok": True, "filename": "y.pptx", "file": "d/y.pptx",
                                             "outDir": "d", "pages": 2}) as m:
            r = self.ws.api_slide_deck(self.ws.SlideDeckReq(
                lesson={"title": "t", "slides": DECK}, theme_id="tech-editorial",
                outDir=str(self.out), confirm=True))
        self.assertTrue(r["ok"])
        self.assertTrue(m.called, "没给模板 → 走主题生成那条路")

    def test_deck_requires_content(self):
        r = self.ws.api_slide_deck(self.ws.SlideDeckReq(lesson={"title": "x", "slides": []}, confirm=True))
        self.assertEqual(getattr(r, "status_code", 200), 400)


class EndpointFillTest(_FillCase):
    """走接口（面板点按钮的真实路径）：`POST /api/slide/fill`。"""

    def setUp(self):
        super().setUp()
        import tempfile as _tf

        self._old_data = os.environ.get("EDU_DATA_DIR")
        os.environ["EDU_DATA_DIR"] = _tf.mkdtemp(prefix="edu_fill_data_")
        from edu_agent import web_server as ws
        from edu_agent import capabilities as cap

        self.ws, self.cap = ws, cap
        cap._CACHE.clear()
        self.addCleanup(self._restore_data)

    def _restore_data(self):
        if self._old_data is None:
            os.environ.pop("EDU_DATA_DIR", None)
        else:
            os.environ["EDU_DATA_DIR"] = self._old_data
        self.cap._CACHE.clear()

    def test_endpoint_fills_and_reports_checks(self):
        from pptx import Presentation
        from pptx.util import Inches

        src = Path(self.tmp.name) / "tpl.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = "标题占位"
        prs.save(str(src))
        tid = st.import_pptx(src, name="接口模板")["id"]

        r = self.ws.api_slide_fill(self.ws.SlideFillReq(
            template_id=tid, slides=DECK, title="3.3 幂函数",
            outDir=str(self.out), confirm=True))
        self.assertTrue(r.get("ok"), r.get("error"))
        self.assertTrue(Path(r["file"]).is_file())
        names = {c["name"]: c["ok"] for c in r["checks"]}
        self.assertTrue(names["内容零丢失"])
        self.assertIn("download_url", r)

    def test_endpoint_needs_template_and_content(self):
        r = self.ws.api_slide_fill(self.ws.SlideFillReq(template_id="", slides=DECK, confirm=True))
        self.assertEqual(getattr(r, "status_code", 200), 400)
        r2 = self.ws.api_slide_fill(self.ws.SlideFillReq(template_id=self.tid, slides=[], confirm=True))
        self.assertEqual(getattr(r2, "status_code", 200), 400)


if __name__ == "__main__":
    unittest.main()
