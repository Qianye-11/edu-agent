# -*- coding: utf-8 -*-
"""演示模板库（M4b / docs/19 §9）单测：导入、抽取、退化、软删、红线。

用**仓库里的样例 pptx** 当"学校模板"（11 页、Office 默认配色），所以要验的是：
1. 配色/字体/字号真的从模板里读出来，且夹进我们的规范区间（越界要记原因）；
2. 槽位抽得出（角色 + 容量 + 几何）；
3. 抽不出来时**退化成只借配色**，不报错、不硬造；
4. **只写模板库**（原稿只读）；内置模板不可删；
5. `EDU_PPT_TEMPLATES` 语义 = **模板库根**（内置在 `_builtin/`）。
"""

from __future__ import annotations

import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

_PROJ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(_PROJ / "src"))

from edu_agent import slide_template as st        # noqa: E402
from edu_agent import ppt_tokens as tk            # noqa: E402

SAMPLE_PPTX = _PROJ / "content" / "plans" / "完整课时模板（默认）.pptx"


class _LibCase(unittest.TestCase):
    """每个用例一个临时模板库（内置 3 套复制进去），别动仓库里的真实模板库。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="edu_tpl_")
        self._old = os.environ.get("EDU_PPT_TEMPLATES")
        os.environ["EDU_PPT_TEMPLATES"] = self.tmp.name
        builtin = Path(self.tmp.name) / "_builtin"
        builtin.mkdir(parents=True)
        for f in (_PROJ / "content" / "ppt_templates" / "_builtin").glob("*.json"):
            shutil.copy(f, builtin / f.name)
        self.addCleanup(self._restore)

    def _restore(self):
        if self._old is None:
            os.environ.pop("EDU_PPT_TEMPLATES", None)
        else:
            os.environ["EDU_PPT_TEMPLATES"] = self._old
        self.tmp.cleanup()


class LibraryTest(_LibCase):
    def test_env_var_means_library_root(self):
        self.assertEqual(tk.templates_root(), Path(self.tmp.name))
        self.assertEqual(tk.builtin_dir(), Path(self.tmp.name) / "_builtin")
        self.assertEqual(tk.available(), ["hand-drawn", "illustration", "tech-editorial"])

    def test_list_builtin_and_imported(self):
        items = st.list_templates()
        self.assertEqual([t["kind"] for t in items], ["json", "json", "json"])
        st.import_pptx(SAMPLE_PPTX, name="学校模板")
        items = st.list_templates()
        self.assertEqual(len(items), 4)
        imported = [t for t in items if t["kind"] == "pptx"][0]
        self.assertEqual(imported["name"], "学校模板")
        self.assertGreater(imported["pages"], 5)
        self.assertTrue(imported["ok"])

    def test_index_file_is_written(self):
        st.import_pptx(SAMPLE_PPTX, name="学校模板")
        self.assertTrue(st.index_path().is_file())
        self.assertIn("学校模板", st.index_path().read_text(encoding="utf-8"))


class ImportTest(_LibCase):
    @unittest.skipUnless(SAMPLE_PPTX.is_file(), "需要仓库样例 pptx")
    def test_theme_extracted_from_template(self):
        r = st.import_pptx(SAMPLE_PPTX, name="学校模板")
        self.assertTrue(r["ok"], r)
        tokens = r["theme"]
        self.assertEqual(tokens["palette"]["text"].startswith("#"), True)
        self.assertEqual(tokens["palette"]["accent"].startswith("#"), True)
        self.assertEqual(tokens["canvas"]["bg"].startswith("#"), True)
        self.assertTrue(tokens["fonts"]["ea"])                   # 中文字体一定有兜底
        self.assertEqual(tk.validate(tokens), [], "抽出来的 tokens 必须自己合法")
        # 字号夹进规范区间：要么本来就合法，要么记了原因
        for key, (lo, hi) in {k: tokens["ranges"][k] for k in ("cover", "title", "body", "caption")}.items():
            self.assertGreaterEqual(tokens["type"][key], lo)
            self.assertLessEqual(tokens["type"][key], hi)

    @unittest.skipUnless(SAMPLE_PPTX.is_file(), "需要仓库样例 pptx")
    def test_slots_extracted(self):
        r = st.import_pptx(SAMPLE_PPTX, name="学校模板")
        self.assertTrue(r["slots_ok"], r.get("degraded"))
        slots = st._read_json(Path(r["dir"]) / "slots.json")     # noqa: SLF001
        self.assertGreater(len(slots["slides"]), 5)
        first = slots["slides"][0]
        self.assertTrue(first["slots"], "至少一个文本槽")
        self.assertTrue(all({"role", "capacity", "geo"} <= set(s) for s in first["slots"]))

    @unittest.skipUnless(SAMPLE_PPTX.is_file(), "需要仓库样例 pptx")
    def test_original_is_read_only_copy(self):
        src_mtime = SAMPLE_PPTX.stat().st_mtime
        r = st.import_pptx(SAMPLE_PPTX, name="学校模板")
        self.assertEqual(SAMPLE_PPTX.stat().st_mtime, src_mtime, "原文件不该被碰")
        self.assertTrue((Path(r["dir"]) / "template.pptx").is_file())

    def test_broken_pptx_degrades(self):
        bad = Path(self.tmp.name) / "bad.pptx"
        bad.write_bytes(b"not a pptx at all")
        r = st.import_pptx(bad, name="坏模板")
        self.assertTrue(r["ok"], "坏包也要能收下（退化成只借配色）")
        self.assertFalse(r["slots_ok"])
        self.assertTrue(any("槽" in d or "配色" in d for d in r["degraded"]), r["degraded"])
        self.assertTrue(st.theme_for(r["id"]), "退化后仍要能渲染（用内置兜底）")

    def test_import_does_not_touch_other_dirs(self):
        watch = (_PROJ / "content" / "plans", _PROJ / "data" / "tmp")

        def snap() -> set:
            out = set()
            for d in watch:
                if d.exists():
                    out |= {str(p) for p in d.rglob("*") if p.is_file()}
            return out

        before = snap()
        st.import_pptx(SAMPLE_PPTX, name="学校模板")
        self.assertEqual(before, snap(), "导入只许写模板库")


class ThemeForTest(_LibCase):
    def test_builtin_id_still_works(self):
        self.assertEqual(st.theme_for("tech-editorial")["id"], "tech-editorial")

    def test_unknown_id_is_none(self):
        self.assertIsNone(st.theme_for("不存在的模板"))
        self.assertIsNone(st.theme_for(""))

    @unittest.skipUnless(SAMPLE_PPTX.is_file(), "需要仓库样例 pptx")
    def test_imported_theme_renders_a_deck(self):
        from edu_agent import slide_deck

        r = st.import_pptx(SAMPLE_PPTX, name="学校模板")
        tokens = st.theme_for(r["id"])
        out = slide_deck.build_deck("", slide_deck.SAMPLE_DECK, title="3.3 幂函数", tokens=tokens)
        self.assertEqual(len(out["pages"]), len(slide_deck.SAMPLE_DECK))
        self.assertIn("<svg", out["pages"][0])


class DeleteTest(_LibCase):
    @unittest.skipUnless(SAMPLE_PPTX.is_file(), "需要仓库样例 pptx")
    def test_soft_delete_moves_to_trash(self):
        r = st.import_pptx(SAMPLE_PPTX, name="学校模板")
        d = st.delete_template(r["id"])
        self.assertTrue(d["ok"])
        self.assertIn("_trash", d["moved_to"])
        self.assertIsNone(st.get(r["id"]))
        self.assertTrue(Path(d["moved_to"], "template.pptx").is_file(), "软删 = 还能找回来")

    def test_builtin_cannot_be_deleted(self):
        r = st.delete_template("tech-editorial")
        self.assertFalse(r["ok"])
        self.assertIn("内置", r["error"])


class DesignBakedRolesTest(_LibCase):
    """角色判定（R7 改）：设计稿模板（**没有占位符**）怎么分标题/正文/装饰。

    这一轮老师的真实模板（12 页、10 种结构、零占位符）暴露了三个坑，
    每个坑都在这里钉一条：
      1. 标题按"面积最大"选 → 4 卡片页上整段正文的框（264×160）比真标题框（400×100）还大，
         标题与正文对调 → 正文被压成几个字塞进小框、真标题框空着（老师："乱"+"正文没填进去"）；
      2. 设计标签只看"≤200×60" → 图形游戏页的「圆」58×68、「矩形」229×126 漏网，
         正文被填进图形标签里；
      3. 容量按设计字号算 → 标题框 66.7px 只算得出 5 个字，40 字的正文被压成 3 个字。
    """

    def _card_template(self) -> Path:
        """仿老师的内容页：大字号标题 + 4 组「小标题（24px）+ 正文段落（16px）」+ 01~04 标签 + 一个更大的页脚标语。"""
        from pptx import Presentation
        from pptx.util import Emu, Pt

        def px(v: float) -> Emu:
            return Emu(int(v * 9525))

        src = Path(self.tmp.name) / "design-baked.pptx"
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[6])                 # 空白版式：全靠文本框
        t = s.shapes.add_textbox(px(440), px(7), px(400), px(100))     # 真标题：字号最大
        t.text_frame.text = "图形的奥秘"
        t.text_frame.paragraphs[0].runs[0].font.size = Pt(50)          # 66.7px
        for i in range(4):                                            # 4 张卡片
            x = 29 + i * 319
            h = s.shapes.add_textbox(px(x), px(331), px(264), px(36))
            h.text_frame.text = f"卡片标题{i + 1}"
            h.text_frame.paragraphs[0].runs[0].font.size = Pt(18)      # 24px
            b = s.shapes.add_textbox(px(x), px(380), px(264), px(160))
            b.text_frame.text = "卡片正文示例文字" * 6
            b.text_frame.paragraphs[0].runs[0].font.size = Pt(12)      # 16px
            n = s.shapes.add_textbox(px(x), px(224), px(85), px(95))
            n.text_frame.text = f"0{i + 1}"                            # 装饰标签
            n.text_frame.paragraphs[0].runs[0].font.size = Pt(36)
        foot = s.shapes.add_textbox(px(23), px(616), px(600), px(108))
        foot.text_frame.text = "争当图形小侦探！"                        # 页脚标语：字号比标题还大
        foot.text_frame.paragraphs[0].runs[0].font.size = Pt(50)
        prs.save(str(src))
        return src

    def _graphic_template(self) -> Path:
        """仿图形游戏页：4 个图形标签（框不大但也不小）+ 一个装不下一句话的小框。"""
        from pptx import Presentation
        from pptx.util import Emu, Pt

        def px(v: float) -> Emu:
            return Emu(int(v * 9525))

        src = Path(self.tmp.name) / "graphic.pptx"
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[6])
        title = s.shapes.add_textbox(px(25), px(19), px(733), px(51))
        title.text_frame.text = "趣味游戏：图形寻踪"
        title.text_frame.paragraphs[0].runs[0].font.size = Pt(25)      # 33.3px
        for label, x, y, w, h in (("圆", 87, 139, 58, 68), ("矩形", 277, 127, 229, 126),
                                  ("正方形", 766, 22, 70, 184), ("三角形", 1032, 302, 72, 184)):
            sh = s.shapes.add_textbox(px(x), px(y), px(w), px(h))
            sh.text_frame.text = label
            sh.text_frame.paragraphs[0].runs[0].font.size = Pt(36)
        body = s.shapes.add_textbox(px(238), px(70), px(371), px(262))  # 真正能写字的框
        body.text_frame.text = ""
        small = s.shapes.add_textbox(px(1171), px(645), px(85), px(24))  # 「小小探险家」那种小标签
        small.text_frame.text = "小小探险家"
        small.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
        prs.save(str(src))
        return src

    def _roles(self, tid: str) -> list[dict]:
        return st.slots_from_pptx(st.root() / tid / "template.pptx")["slides"][0]["slots"]

    def test_title_is_picked_by_font_size_not_area(self):
        tid = st.import_pptx(self._card_template(), name="设计稿模板")["id"]
        titles = [s for s in self._roles(tid) if s["role"] == "title"]
        self.assertEqual([s["sample"] for s in titles], ["图形的奥秘"],
                         "最大面积的框是正文段落 → 标题必须按字号选")

    def test_footer_callout_with_bigger_font_is_not_title(self):
        tid = st.import_pptx(self._card_template(), name="设计稿模板")["id"]
        foot = [s for s in self._roles(tid) if s["sample"] == "争当图形小侦探！"][0]
        self.assertEqual(foot["role"], "decoration", "贴着页面底边的标语不是标题（也不该被填正文）")

    def test_small_graphic_labels_are_decoration(self):
        tid = st.import_pptx(self._graphic_template(), name="图形页模板")["id"]
        roles = {s["sample"]: s["role"] for s in self._roles(tid)}
        for label in ("圆", "矩形", "正方形", "三角形"):
            with self.subTest(label=label):
                self.assertEqual(roles.get(label), "decoration", "图形页的图形名是设计标签")
        self.assertEqual(roles.get("小小探险家"), "decoration", "装不下一句话的小框（容量 5）不算正文")
        self.assertEqual(roles.get("趣味游戏：图形寻踪"), "title")

    def test_capacity_is_measured_at_the_smallest_allowed_font(self):
        tid = st.import_pptx(self._card_template(), name="设计稿模板")["id"]
        title = [s for s in self._roles(tid) if s["role"] == "title"][0]
        self.assertGreater(title["capacity"], title["design_capacity"],
                           "容量按可缩到的最小字号算 → 比设计字号下的容量大（否则 40 字被压成 5 字）")

    def test_placeholder_title_wins_over_short_text_rule(self):
        """有占位符的模板：标题占位符里的示例文字常常又长又两行，不能被"标题要短"排除掉。"""
        from pptx import Presentation
        from pptx.util import Inches, Pt

        src = Path(self.tmp.name) / "ph.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        tf = slide.shapes.title.text_frame
        tf.paragraphs[0].add_run().text = "点击此处添加标题（模板原名）"
        tf.paragraphs[0].runs[0].font.size = Pt(32)
        tf.add_paragraph().add_run().text = "（占位第二行）"
        body = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(0.6))
        body.text_frame.text = "模板自己的装饰文字"
        prs.save(str(src))
        tid = st.import_pptx(src, name="占位符模板")["id"]
        titles = [s for s in self._roles(tid) if s["role"] == "title"]
        self.assertEqual(len(titles), 1)
        self.assertEqual(titles[0]["shape_index"], 0, "标题占位符就是标题槽")
        self.assertEqual([s["role"] for s in self._roles(tid) if s["shape_index"] == 1], ["body"])


if __name__ == "__main__":
    unittest.main()
    """槽位对齐（R6 之后）：老师在一个结构上标一次，**所有同结构页都跟着对**。"""

    def _two_layout_template(self) -> Path:
        from pptx import Presentation
        from pptx.util import Inches

        src = Path(self.tmp.name) / "align-tpl.pptx"
        prs = Presentation()
        for _ in range(2):                       # 两页**同结构**（内容页）
            s = prs.slides.add_slide(prs.slide_layouts[5])
            s.shapes.title.text_frame.text = "示例标题文字"
            s.shapes.add_textbox(Inches(1), Inches(3), Inches(4), Inches(2)).text_frame.text = "示例正文内容文字"
        prs.save(str(src))
        return src

    def test_signature_groups_same_structure(self):
        tid = st.import_pptx(self._two_layout_template(), name="对齐模板")["id"]
        o = st.align_overview(tid)
        self.assertTrue(o["ok"], o.get("error"))
        self.assertEqual(len(o["structures"]), 1, "两页同结构 → 一个结构")
        self.assertEqual(o["structures"][0]["pages"], 2)

    def test_align_override_applies_to_all_same_structure_pages(self):
        tid = st.import_pptx(self._two_layout_template(), name="对齐模板")["id"]
        sig = st.align_overview(tid)["structures"][0]["signature"]
        slots = st.slots_from_pptx(st.root() / tid / "template.pptx")["slides"]
        body_idx = [s["shape_index"] for s in slots[0]["slots"] if s["role"] == "body"][0]
        st.save_align(tid, sig, {str(body_idx): "decoration"})          # 老师把"正文"标成装饰
        after = st.slots_from_pptx(st.root() / tid / "template.pptx")["slides"]
        for page in after:
            with self.subTest(page=page["index"]):
                got = [s["role"] for s in page["slots"] if s["shape_index"] == body_idx][0]
                self.assertEqual(got, "decoration", "同结构页都要跟着改")
        self.assertTrue(st.load_align(tid)["pages"][sig])

    def test_bad_role_is_ignored(self):
        tid = st.import_pptx(self._two_layout_template(), name="对齐模板")["id"]
        sig = st.align_overview(tid)["structures"][0]["signature"]
        st.save_align(tid, sig, {"0": "乱写的角色"})
        self.assertEqual(st.load_align(tid)["pages"][sig], {})
