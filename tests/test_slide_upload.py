# -*- coding: utf-8 -*-
"""插入图片（R22）：本地上传落 `data/uploads/`、缩一版大图、当对象渲染、导出带图。

三条要钉住的：

1. **上传入口要严**：空文件/非图片/超大 一律拒；只写 `data/uploads/`（`content/plans` 一个字不写）；
2. **大图先缩一版**（走 ppt-master 的 `image_treat.py --fit`，**非破坏性**：原图必须还在）；
3. **导出真的带图**：预览与导出的同一份 SVG 里是内联 `data:` 图，pptx 里有图片部件。
"""

from __future__ import annotations

import asyncio
import os
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

_PROJ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(_PROJ / "src"))

from edu_agent import ppt_svg                            # noqa: E402
from edu_agent import slide_deck as sd                   # noqa: E402
from edu_agent import slide_image as si                  # noqa: E402

THEME = "tech-editorial"


def _png(w: int = 40, h: int = 30) -> bytes:
    """造一张真 PNG（用 Pillow，别手搓字节 —— 免得"测试通过、真图打不开"）。"""
    import io

    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (w, h), (220, 230, 245)).save(buf, format="PNG")
    return buf.getvalue()


def _jpeg(w: int, h: int) -> bytes:
    import io

    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (w, h), (200, 200, 200)).save(buf, format="JPEG", quality=70)
    return buf.getvalue()


class _Case(unittest.TestCase):
    def setUp(self):
        self.data = Path(tempfile.mkdtemp(prefix="edu_up_"))
        self._old = os.environ.get("EDU_DATA_DIR")
        os.environ["EDU_DATA_DIR"] = str(self.data)
        self.addCleanup(self._restore)

    def _restore(self):
        if self._old is None:
            os.environ.pop("EDU_DATA_DIR", None)
        else:
            os.environ["EDU_DATA_DIR"] = self._old


class SaveUploadTest(_Case):
    def test_saves_into_uploads_only(self):
        got = si.save_upload(_png(), "老师的图.png", "sess-1")
        self.assertTrue(got["ok"], got)
        self.assertTrue(got["path"].startswith("sess-1/"), got["path"])
        self.assertTrue(Path(got["file"]).is_file())
        self.assertEqual(got["width"], 40)
        self.assertEqual(got["height"], 30)
        plans = _PROJ / "content" / "plans"
        before = sorted(p.name for p in plans.glob("*")) if plans.is_dir() else []
        si.save_upload(_png(), "再来一张.png", "sess-1")
        after = sorted(p.name for p in plans.glob("*")) if plans.is_dir() else []
        self.assertEqual(before, after, "content/plans 一个字都不能动")

    def test_rejects_bad_input(self):
        self.assertFalse(si.save_upload(b"", "empty.png", "s")["ok"])
        self.assertIn("太大", si.save_upload(b"x" * (si.MAX_UPLOAD_BYTES + 1), "big.png", "s")["error"])
        # 纯文本伪装成 .png：按文件头兜底 → 拒
        bad = si.save_upload(b"hello world" * 20, "fake.png", "s")
        self.assertFalse(bad["ok"], bad)
        self.assertIn("不是图片", bad["error"])

    def test_sniffs_type_when_name_lies(self):
        """名字是 .txt 但内容是真 PNG：按文件头认出来（老师从聊天软件存下来的文件常这样）。"""
        got = si.save_upload(_png(), "屏幕截图.txt", "s")
        self.assertTrue(got["ok"], got)
        self.assertTrue(got["path"].endswith(".png"), got["path"])

    def test_probe_reads_size(self):
        got = si.save_upload(_jpeg(120, 60), "photo.jpg", "s")
        info = si.probe(Path(got["file"]))
        self.assertEqual((info["width"], info["height"]), (120, 60))
        self.assertEqual(si.probe(self.data / "nope.png")["width"], 0)


class OptimizeTest(_Case):
    def test_small_image_is_left_alone(self):
        got = si.save_upload(_png(320, 200), "small.png", "s")
        opt = si.optimize(got["path"])
        self.assertTrue(opt["ok"])
        self.assertEqual(opt["path"], got["path"], "小图不该被动")
        self.assertEqual(opt["treated"], "")

    def test_big_image_gets_a_smaller_derivative_original_kept(self):
        """大图：要么真的缩出一版（treated=fit），要么如实说缩不了 —— 不许假装。"""
        big = _jpeg(2400, 1600)
        got = si.save_upload(big, "big.jpg", "s")
        opt = si.optimize(got["path"])
        self.assertTrue(opt["ok"], opt)
        self.assertTrue(Path(got["file"]).is_file(), "原图必须还在（非破坏性）")
        if opt["treated"] == "fit":
            self.assertNotEqual(opt["path"], got["path"])
            self.assertTrue((si.uploads_root() / opt["path"]).is_file())
            self.assertLessEqual(max(opt["info"]["width"], opt["info"]["height"]), si.FIT_MAX_SIDE)
        else:
            self.assertTrue(opt.get("note"), "没缩成也得说明原因")

    def test_optimize_rejects_path_escape(self):
        self.assertFalse(si.optimize("../../etc/passwd")["ok"])


class ImageObjectTest(_Case):
    def _deck(self, src: str):
        return [{"layout": "bullets", "name": "一页", "blocks": [{"text": "要点"}],
                 "objects": [{"kind": "image", "id": "o1", "src": src, "x": 100, "y": 120,
                              "w": 400, "h": 240, "fit": "cover"}]}]

    def test_clean_objects_keeps_and_validates_images(self):
        got = sd.clean_objects([{"kind": "image", "src": "s1/a.png", "fit": "contain"},
                                {"kind": "image", "src": "../secret.png"},
                                {"kind": "image", "src": "noext"},
                                {"kind": "image"}])
        self.assertEqual(len(got), 1, got)
        self.assertEqual(got[0]["src"], "s1/a.png")
        self.assertEqual(got[0]["fit"], "contain")

    def test_render_inlines_the_picture(self):
        got = si.save_upload(_png(200, 100), "pic.png", "s1")
        deck = sd.build_deck(THEME, self._deck(got["path"]), "x", annotate=True)
        svg = deck["pages"][0]
        self.assertIn("<image ", svg)
        self.assertIn('href="data:image/png;base64,', svg)
        self.assertIn('preserveAspectRatio="xMidYMid slice"', svg)
        box = [b for b in deck["boxes"][0] if b["slot"] == "obj:o1"][0]
        self.assertEqual(box["kind"], "obj-image")
        self.assertEqual((box["x"], box["y"], box["w"], box["h"]), (100.0, 120.0, 400.0, 240.0))

    def test_fit_contain_switches_aspect_rule(self):
        got = si.save_upload(_png(200, 100), "pic.png", "s1")
        deck = self._deck(got["path"])
        deck[0]["objects"][0]["fit"] = "contain"
        svg = sd.build_deck(THEME, deck, "x")["pages"][0]
        self.assertIn('preserveAspectRatio="xMidYMid meet"', svg)

    def test_missing_file_renders_nothing_but_keeps_box(self):
        deck = sd.build_deck(THEME, self._deck("s1/ghost.png"), "x", annotate=True)
        self.assertNotIn("<image ", deck["pages"][0], "取不到的图不该留个破图")
        self.assertTrue([b for b in deck["boxes"][0] if b["slot"] == "obj:o1"], "盒子还在（能选中删掉）")

    def test_export_path_draws_picture(self):
        got = si.save_upload(_png(200, 100), "pic.png", "s1")
        page = ppt_svg.render_deck(ppt_svg.tk.load(THEME), self._deck(got["path"]))[0]
        self.assertIn("<image ", page)
        self.assertIn("data:image/png;base64,", page)


class UploadEndpointTest(_Case):
    class _Fake:
        def __init__(self, name: str, data: bytes):
            self.filename = name
            self._data = data

        async def read(self):
            return self._data

    def test_upload_endpoint(self):
        from edu_agent import web_server as ws

        sid = "up-1"
        ws._store.set_session_doc(sid, {"slides": [{"layout": "bullets", "name": "x",
                                                    "blocks": [{"text": "一"}]}]})
        got = asyncio.run(ws.api_slide_image_upload(file=self._Fake("图.png", _png(64, 48)),
                                                   session_id=sid, treat="0"))
        self.assertTrue(got["ok"], got)
        self.assertTrue(got["path"].startswith(sid + "/"), got["path"])
        self.assertEqual((got["width"], got["height"]), (64, 48))
        self.assertNotIn("file", got, "别把服务器绝对路径回给前端")

        bad = asyncio.run(ws.api_slide_image_upload(file=self._Fake("x.png", b"not an image"),
                                                   session_id=sid, treat="0"))
        self.assertEqual(getattr(bad, "status_code", 200), 400)

    def test_object_add_with_image_keeps_aspect(self):
        from edu_agent import web_server as ws

        sid = "up-2"
        slides = [{"layout": "bullets", "name": "x", "blocks": [{"text": "一"}]}]
        ws._store.set_session_doc(sid, {"slides": slides})
        got = asyncio.run(ws.api_slide_image_upload(file=self._Fake("wide.png", _png(400, 100)),
                                                   session_id=sid, treat="0"))
        add = ws.api_slide_object(ws.SlideObjectReq(session_id=sid, index=0, action="add",
                                                    kind="image", src=got["path"], slides=slides))
        self.assertTrue(add["ok"], add)
        obj = add["objects"][0]
        self.assertEqual(obj["kind"], "image")
        self.assertEqual(obj["src"], got["path"])
        self.assertAlmostEqual(obj["w"] / obj["h"], 4.0, places=1, msg="按原比例摆框，不拉变形")


class ExportImageTest(_Case):
    def setUp(self):
        super().setUp()
        try:
            from edu_agent import ppt_polish
            self.ready = bool(ppt_polish.status().get("intake_ready"))
        except Exception:                                     # noqa: BLE001
            self.ready = False

    def test_export_carries_picture(self):
        if not self.ready:
            self.skipTest("没装 ppt-master")
        got = si.save_upload(_png(240, 120), "pic.png", "s1")
        deck = [{"layout": "cover", "title": "带图的稿"},
                {"layout": "bullets", "name": "一页", "blocks": [{"text": "要点"}],
                 "objects": [{"kind": "image", "id": "o1", "src": got["path"],
                              "x": 120, "y": 200, "w": 400, "h": 200}]}]
        out = sd.export_deck_pptx(THEME, deck, "带图验收")
        self.assertTrue(out.get("ok"), out.get("error"))
        with zipfile.ZipFile(out["file"]) as z:
            media = [n for n in z.namelist() if n.startswith("ppt/media/")]
        self.assertTrue(media, "产物里没有图片部件")
        from pptx import Presentation

        prs = Presentation(out["file"])
        pics = [s for s in prs.slides[1].shapes if str(s.shape_type).startswith("PICTURE")]
        self.assertTrue(pics, [(s.name, str(s.shape_type)) for s in prs.slides[1].shapes])
        self.assertAlmostEqual(int(pics[0].width) / int(pics[0].height), 2.0, places=1)


if __name__ == "__main__":
    unittest.main()
