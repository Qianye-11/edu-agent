"""把教案中心的可编辑 HTML 直接导出为 DOCX，不依赖 WPS。"""
from __future__ import annotations

import base64
from io import BytesIO
import re
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "content" / "plans"


class Node:
    def __init__(self, tag: str = "root", attrs: dict[str, str] | None = None):
        self.tag = tag
        self.attrs = attrs or {}
        self.children: list[Node | str] = []


class TreeBuilder(HTMLParser):
    VOID = {"br", "img", "hr", "meta", "link", "input"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = Node()
        self.stack = [self.root]

    def handle_starttag(self, tag: str, attrs):
        node = Node(tag.lower(), {str(k).lower(): str(v or "") for k, v in attrs})
        self.stack[-1].children.append(node)
        if node.tag not in self.VOID:
            self.stack.append(node)

    def handle_startendtag(self, tag: str, attrs):
        self.handle_starttag(tag, attrs)
        if tag.lower() not in self.VOID:
            self.handle_endtag(tag)

    def handle_endtag(self, tag: str):
        tag = tag.lower()
        for index in range(len(self.stack) - 1, 0, -1):
            if self.stack[index].tag == tag:
                del self.stack[index:]
                break

    def handle_data(self, data: str):
        if data:
            self.stack[-1].children.append(data)


def _font(run, name: str = "宋体", size: float | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)


def _style_value(style: str, name: str) -> str:
    hit = re.search(rf"(?:^|;)\s*{re.escape(name)}\s*:\s*([^;]+)", style, flags=re.I)
    return hit.group(1).strip() if hit else ""


def _inline(paragraph, node: Node | str, bold: bool = False, italic: bool = False,
            underline: bool = False) -> None:
    if isinstance(node, str):
        if not node:
            return
        run = paragraph.add_run(node)
        run.bold, run.italic, run.underline = bold, italic, underline
        _font(run)
        return
    if node.tag == "br":
        paragraph.add_run().add_break()
        return
    if node.tag == "img":
        # 编辑画布中的本地图片以 data URL 保存；导出时保留，而不是退化为“[图片]”。
        src = str(node.attrs.get("src") or "")
        hit = re.fullmatch(r"data:image/(?:png|jpeg|jpg|gif|webp);base64,([A-Za-z0-9+/=\s]+)", src,
                           flags=re.I)
        if hit:
            try:
                raw = base64.b64decode(re.sub(r"\s+", "", hit.group(1)), validate=True)
                if raw and len(raw) <= 8 * 1024 * 1024:
                    paragraph.add_run().add_picture(BytesIO(raw))
                    return
            except Exception:  # noqa: BLE001 - keep exporting surrounding document on bad image data
                pass
        run = paragraph.add_run("[图片]")
        run.italic = True
        _font(run)
        return
    next_bold = bold or node.tag in {"b", "strong"}
    next_italic = italic or node.tag in {"i", "em"}
    next_underline = underline or node.tag == "u"
    start = len(paragraph.runs)
    for child in node.children:
        _inline(paragraph, child, next_bold, next_italic, next_underline)
    style = node.attrs.get("style", "")
    for run in paragraph.runs[start:]:
        family = _style_value(style, "font-family").strip(" '\"").split(",")[0] or "宋体"
        size_text = _style_value(style, "font-size")
        size_hit = re.search(r"[0-9.]+", size_text)
        _font(run, family, float(size_hit.group()) if size_hit else None)
        color = _style_value(style, "color").lstrip("#")
        if re.fullmatch(r"[0-9a-fA-F]{6}", color):
            run.font.color.rgb = RGBColor.from_string(color.upper())


def _text(node: Node) -> str:
    out: list[str] = []
    for child in node.children:
        if isinstance(child, str):
            out.append(child)
        elif child.tag == "br":
            out.append("\n")
        else:
            out.append(_text(child))
    return "".join(out).strip()


def _paragraph(document_or_cell, node: Node, style: str | None = None):
    paragraph = document_or_cell.add_paragraph(style=style) if style else document_or_cell.add_paragraph()
    for child in node.children:
        _inline(paragraph, child)
    align = _style_value(node.attrs.get("style", ""), "text-align").lower()
    if node.tag == "h1" or align == "center":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if node.tag in {"h1", "h2", "h3"}:
        sizes = {"h1": 18, "h2": 15, "h3": 13}
        for run in paragraph.runs:
            run.bold = True
            _font(run, "黑体" if node.tag != "h1" else "宋体", sizes[node.tag])
    paragraph.paragraph_format.space_after = Pt(10 if node.tag == "h1" else 4)
    paragraph.paragraph_format.line_spacing = 1.35
    return paragraph


def _cell_content(cell, node: Node) -> None:
    cell.text = ""
    blocks = [x for x in node.children if isinstance(x, Node) and x.tag in {"p", "div", "h1", "h2", "h3", "ul", "ol"}]
    if not blocks:
        p = cell.paragraphs[0]
        for child in node.children:
            _inline(p, child)
        return
    first = True
    for block in blocks:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        if block.tag in {"ul", "ol"}:
            for li in [x for x in block.children if isinstance(x, Node) and x.tag == "li"]:
                target = p if p.text == "" else cell.add_paragraph()
                target.style = "List Bullet" if block.tag == "ul" else "List Number"
                for child in li.children:
                    _inline(target, child)
        else:
            for child in block.children:
                _inline(p, child)
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.2


def _set_cell_width(cell, width_cm: float) -> None:
    cell.width = Cm(width_cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(round(width_cm / 2.54 * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def _set_cell_margins(cell, *, top: int = 80, start: int = 110,
                      bottom: int = 80, end: int = 110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_layout(table, column_widths: list[float]) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is not None:
        tbl_w.set(qn("w:w"), "5000")
        tbl_w.set(qn("w:type"), "pct")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    for index, width_cm in enumerate(column_widths):
        table.columns[index].width = Cm(width_cm)
        grid_col = table._tbl.tblGrid.gridCol_lst[index]
        grid_col.set(qn("w:w"), str(round(width_cm / 2.54 * 1440)))


def _set_row_min_height(row, height_cm: float, *, allow_split: bool = False) -> None:
    row.height = Cm(height_cm)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    if not allow_split:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))


def _lesson_row_height(label: str, row_index: int) -> float:
    label = re.sub(r"\s+", "", label)
    heights = {
        "课程分析": 0.85,
        "教学目标": 2.25,
        "重点难点": 0.72,
        "课前准备": 0.72,
        # 网页模板中 .lp-process-label 的高度为 430px，约等于 11.4cm。
        "教学过程": 11.4,
        "板书与作业": 0.95,
        "课后反思": 1.15,
    }
    return heights.get(label, 0.66 if row_index < 2 else 0.72)


def _table(document, node: Node) -> None:
    rows = [x for x in node.children if isinstance(x, Node) and x.tag == "tr"]
    if not rows:
        for section in [x for x in node.children if isinstance(x, Node) and x.tag in {"thead", "tbody", "tfoot"}]:
            rows.extend(x for x in section.children if isinstance(x, Node) and x.tag == "tr")
    if not rows:
        return
    cell_rows = [[x for x in row.children if isinstance(x, Node) and x.tag in {"td", "th"}] for row in rows]
    width = max((sum(max(1, int(c.attrs.get("colspan", "1") or "1")) for c in cells) for cells in cell_rows), default=1)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    # 规范教案模板的网页列宽比例约为 92:125，并以三组“标签 + 内容”排布。
    if width == 6:
        column_widths = [2.42, 3.30, 2.42, 3.30, 2.42, 3.30]
    else:
        usable_width = 17.16
        column_widths = [usable_width / width] * width
    _set_table_layout(table, column_widths)
    for ri, cells in enumerate(cell_rows):
        column = 0
        for source in cells:
            span = max(1, int(source.attrs.get("colspan", "1") or "1"))
            target = table.cell(ri, min(column, width - 1))
            if span > 1 and column + span - 1 < width:
                target = target.merge(table.cell(ri, column + span - 1))
            _set_cell_width(target, sum(column_widths[column:column + span]))
            _set_cell_margins(target)
            _cell_content(target, source)
            is_label = source.tag == "th" or "lp-label" in source.attrs.get("class", "")
            if is_label:
                target.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for p in target.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.bold = True
            else:
                target.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            column += span
        label = _text(cells[0]) if cells else ""
        compact_label = re.sub(r"\s+", "", label)
        _set_row_min_height(
            table.rows[ri],
            _lesson_row_height(label, ri),
            # 生成完整教案后该行可能跨页，允许拆分可避免长内容被裁切。
            allow_split=compact_label == "教学过程",
        )


def export_html(source_html: str, title: str, out_dir: Path | None = None) -> Path:
    if not str(source_html or "").strip():
        raise ValueError("教案正文不能为空")
    if len(source_html.encode("utf-8")) > 2_000_000:
        raise ValueError("教案内容超过 2MB")
    parser = TreeBuilder()
    parser.feed(source_html)
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    normal = document.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    for node in parser.root.children:
        if isinstance(node, str):
            if node.strip():
                p = document.add_paragraph(node.strip())
                for run in p.runs:
                    _font(run)
        elif node.tag == "div" and "lesson-page-break" in node.attrs.get("class", "").split():
            document.add_page_break()
        elif node.tag in {"h1", "h2", "h3", "p", "div"}:
            _paragraph(document, node)
        elif node.tag == "table":
            _table(document, node)
        elif node.tag in {"ul", "ol"}:
            for li in [x for x in node.children if isinstance(x, Node) and x.tag == "li"]:
                _paragraph(document, li, "List Bullet" if node.tag == "ul" else "List Number")
        elif node.tag == "img":
            document.add_paragraph("[图片]")
    output = out_dir or OUT_DIR
    output.mkdir(parents=True, exist_ok=True)
    safe = re.sub(r'[\\/:*?"<>|]+', "_", title or "教案").strip(" .")[:70] or "教案"
    path = output / f"{safe}.docx"
    document.save(path)
    return path
