"""幻灯片渲染收口：主题 + 会话数据 → 每页 SVG（预览与导出共用**同一份**）。

为什么要有这一层（`docs/22` D13 / `docs/23` M1）：

- 以前面板用 JS 拼幻灯片 HTML、导出用 python-pptx 另写坐标 ⇒ "预览≠导出"永远修不掉；
- 现在只有**一份**布局逻辑（`ppt_svg`），调用方只有两个：
  面板预览（`POST /api/slide/render`，纯内存）与导出（M3，写 `data/tmp/ppt_jobs/<job>/svg_output/`）。

本模块**绝不写盘**：`build_deck()` 只返回字符串，落盘一律由调用方按红线去做（有单测锁）。
"""

from __future__ import annotations

import html
import json
import math
import os
import re
import subprocess
import sys
import time
import zipfile
from pathlib import Path

from . import ppt_svg
from . import ppt_tokens as tk
from . import slide_image

#: 版式 id（与 docs/22 §5 一致；显式 `layout` 只认这些，其余按位置推断）
LAYOUTS = ppt_svg.LAYOUTS

#: 演示用的示例教案（样张工具与审计工具共用一份，避免两处各写一套）
SAMPLE_DECK: list[dict] = [
    {"layout": "cover", "name": "3.3 幂函数", "cover_title": "3.3 幂函数",
     "subtitle": "数学 · 高一 · 新授课", "label": "人教A版 · 必修第一册",
     "meta": "1 课时 · 情境导入 → 概念建构 → 图象探究 → 小结"},
    {"layout": "bullets", "name": "幂函数的概念", "minutes": "12", "label": "环节 2 · 12 分钟",
     "blocks": [{"label": "定义", "text": "形如 y=x^α 的函数，α 为常数"},
                {"label": "关键", "text": "底数是自变量，指数是常数"},
                {"label": "定义域", "text": "随 α 的取值变化，先看指数再定"},
                {"label": "五个常见", "text": "y=x、y=x²、y=x³、y=√x、y=x⁻¹"}]},
    {"layout": "media", "name": "图象与性质", "minutes": "15", "label": "环节 3 · 15 分钟",
     "caption": "在同一坐标系里画出 y=x²，观察它的单调性与奇偶性。",
     "bullets": ["关于 y 轴对称，是偶函数", "在 [0,+∞) 上单调递增"]},
]


# ---------------------------------------------------------------- 主题

def themes() -> list[dict]:
    """内置主题清单（给面板的下拉用）：数值 + 出处 + 字体是否本机可用。"""
    out: list[dict] = []
    for theme_id in tk.available():
        t = tk.load(theme_id)
        problems = tk.validate(t)
        out.append({
            "id": t.get("id") or theme_id,
            "name": t.get("name") or theme_id,
            "source": t.get("source") or "",
            "canvas": t["canvas"],
            "palette": {k: v for k, v in t["palette"].items() if k != "accent_max_ratio"},
            "type": t["type"],
            "fonts_ok": not any("字体" in p for p in problems),
            "fonts_missing": [p for p in problems if "字体" in p],
        })
    return out


#: 首次打开时的默认主题顺序（docs/23 §7：默认"学术蓝"，最正式、说课/公开课都稳）
PREFERRED = ("tech-editorial", "illustration", "hand-drawn")


def default_theme_id() -> str:
    """默认主题：按 `PREFERRED` 取第一个可用的（面板会记住老师选过的那套）。"""
    ids = tk.available()
    for want in PREFERRED:
        if want in ids:
            return want
    return ids[0] if ids else ""


def theme(theme_id: str) -> dict:
    """取一套主题（空值 → 默认主题）。"""
    return tk.load(theme_id or default_theme_id())


# ---------------------------------------------------------------- 会话数据 → 槽位

def _text(v: object) -> str:
    return str(v or "").strip()


def _bullets(slide: dict) -> list[str]:
    raw = slide.get("bullets")
    if not isinstance(raw, list):
        return []
    return [_text(b) for b in raw if _text(b)]


def _blocks(slide: dict) -> list[dict]:
    """`blocks[]`（带标签）优先；否则把 `bullets[]` 当成无标签的块。

    两种写法都收，因为会话历史里是 `bullets`，而模板/手工数据里是 `blocks`。
    """
    raw = slide.get("blocks")
    if isinstance(raw, list) and raw:
        out = []
        for b in raw:
            if isinstance(b, dict):
                text = _text(b.get("text") or b.get("name"))
                if text:
                    out.append({"label": _text(b.get("label")), "text": text})
            elif _text(b):
                out.append({"label": "", "text": _text(b)})
        return out
    return [{"label": "", "text": b} for b in _bullets(slide)]


def _items(slide: dict) -> list[str]:
    """取这一页的"条目文字"（`bullets` 优先，其次 `blocks` 的正文）——用于要点页/图文页/结尾页。"""
    items = _bullets(slide)
    return items or [b["text"] for b in _blocks(slide)]


#: 可编辑槽 → `clean_slides` 字段的映射（**唯一真相源**，R10）。
#: 前端拿到 `boxes[].slot` 后，改动都回传到 `POST /api/slide/slot`，由这里的 `apply_slot` 落字段 ——
#: 两边各写一套映射迟早会漂（改字段名时这里会红，前端只用 slot 字符串，不认字段名）。
_SLOT_RE = re.compile(r"^(?P<base>title|subtitle|meta|label|closing|caption|index|cover_title"
                      r"|block:(?P<bi>\d+)(?::label)?"
                      r"|(?P<side>left|right):(?:heading|bullet:(?P<ci>\d+)))$")


def iter_text_slots(slide: dict) -> list[tuple[str, str]]:
    """一页里**所有可编辑文字**的 `(槽位, 文本)` 清单（R18：查找替换与页内编辑共用这一份）。

    为什么要有它：查找替换如果自己再走一遍字段，迟早和 `apply_slot` 的槽位命名漂开
    （一个改 `blocks`、一个改 `bullets`，最后"找得到却替换不了"）。
    这里只列举**当前真有字的槽**，顺序就是阅读顺序（标题 → 标签 → 正文块 → 两栏）。
    """
    out: list[tuple[str, str]] = []
    title = _text(slide.get("name") or slide.get("title"))
    if title:
        out.append(("title", title))
    for key in ("subtitle", "meta", "label", "caption", "closing", "index", "cover_title"):
        v = _text(slide.get(key))
        if not v:
            continue
        # 封面页的 `cover_title` 与 `name` 常常是同一串字 —— 那是**同一处**文字（渲染时取其一），
        # 报两遍会让老师看到"找到 4 处"而页面上只有 3 处（踩过：替换数也跟着翻倍）。
        if key == "cover_title" and v == title:
            continue
        out.append((key, v))
    for i, b in enumerate(_blocks(slide) or [{"label": "", "text": b} for b in _bullets(slide)]):
        if _text(b.get("label")):
            out.append((f"block:{i}:label", _text(b["label"])))
        if _text(b.get("text")):
            out.append((f"block:{i}", _text(b["text"])))
    for side in ("left", "right"):
        col = slide.get(side) if isinstance(slide.get(side), dict) else {}
        if _text(col.get("heading")):
            out.append((f"{side}:heading", _text(col["heading"])))
        for j, x in enumerate(col.get("bullets") or []):
            if _text(x):
                out.append((f"{side}:bullet:{j}", _text(x)))
    return out


def _find_spans(text: str, query: str, *, case_sensitive: bool = False) -> list[tuple[int, int]]:
    """`query` 在 `text` 里的出现位置（**按字面找，不当正则**：老师输入 `(` 不该报错）。"""
    q, t = str(query or ""), str(text or "")
    if not q or not t:
        return []
    if case_sensitive:
        hay, needle = t, q
    else:
        hay, needle = t.lower(), q.lower()
    spans, i = [], hay.find(needle)
    while i >= 0:
        spans.append((i, i + len(q)))
        i = hay.find(needle, i + len(q))
    return spans


def find_text(slides: list[dict] | None, query: str, *, case_sensitive: bool = False,
              pages: list[int] | None = None) -> list[dict]:
    """跨页找字：返回 `[{index, slot, text, snippet, count}]`（`pages` 给了就只找那几页）。

    备注（`notes`）也算 —— 它是要讲给学生的台词，找不到会很奇怪；但它的槽位名不是版式槽，
    所以 `slot` 给 `"notes"`，前端知道"跳页 + 高亮备注框"而不是"选中某一格"。
    """
    want = set(pages) if pages is not None else None
    hits: list[dict] = []
    for i, s in enumerate(slides or []):
        if not isinstance(s, dict) or (want is not None and i not in want):
            continue
        for slot, text in iter_text_slots(s):
            spans = _find_spans(text, query, case_sensitive=case_sensitive)
            if spans:
                a, b = spans[0]
                hits.append({"index": i, "slot": slot, "text": text, "count": len(spans),
                             "snippet": text[max(0, a - 8):b + 12]})
        note = _text(s.get("notes"))
        if note and _find_spans(note, query, case_sensitive=case_sensitive):
            hits.append({"index": i, "slot": "notes", "text": note, "count": 1,
                         "snippet": note[:40]})
    return hits


def replace_text(slides: list[dict] | None, query: str, replace: str, *,
                 case_sensitive: bool = False, pages: list[int] | None = None,
                 slots: list[str] | None = None) -> dict:
    """跨页替换（**走 `apply_slot` 同一条写路径**）：返回 `{slides, changed, pages, hits}`。

    `slots` 给了就只替换那几个槽（前端"只替换这一处"用）；`replace` 给空串 = 删除匹配。
    改完的 `slides` 交给调用方过一遍 `clean_slides` 再落库（与页内编辑一致）。
    """
    want = set(pages) if pages is not None else None
    only = set(slots) if slots else None
    out = [dict(s) if isinstance(s, dict) else {} for s in (slides or [])]
    changed, touched_pages, hits = 0, [], []
    for i, s in enumerate(out):
        if want is not None and i not in want:
            continue
        for slot, text in iter_text_slots(s):
            if only is not None and slot not in only:
                continue
            spans = _find_spans(text, query, case_sensitive=case_sensitive)
            if not spans:
                continue
            new = text
            for a, b in reversed(spans):
                new = new[:a] + str(replace) + new[b:]
            res = apply_slot(s, slot, new)
            if res.get("ok"):
                changed += len(spans)
                hits.append({"index": i, "slot": slot, "before": text, "after": new,
                             "count": len(spans)})
        note = _text(s.get("notes"))
        if note and (only is None or "notes" in only):
            spans = _find_spans(note, query, case_sensitive=case_sensitive)
            if spans:
                new = note
                for a, b in reversed(spans):
                    new = new[:a] + str(replace) + new[b:]
                s["notes"] = new[:2000]
                changed += len(spans)
                hits.append({"index": i, "slot": "notes", "before": note, "after": new,
                             "count": len(spans)})
        if any(h["index"] == i for h in hits) and i not in touched_pages:
            touched_pages.append(i)
    return {"slides": out, "changed": changed, "pages": touched_pages, "hits": hits}


def apply_slot(slide: dict, slot: str, text: str) -> dict:
    """把"某个槽的新文字"落进一页 slide（返回同一份 dict，改完请再过 `clean_slides`）。

    槽位命名（与 `ppt_svg` 打点、前端 `boxes[].slot` 完全一致）：
      `title` / `subtitle` / `meta` / `label` / `closing` / `caption` /
      `block:<i>` / `block:<i>:label` / `left:heading` / `left:bullet:<i>` / 同 `right:*`

    边界（都不抛异常，返回 `{"ok": False, "error": ...}` 让调用方去报）：
      · 不认识的槽 → 拒绝（不猜）；
      · 越界的块/条目下标 → 补齐到该下标（老师加一块就等于加一个块）；
      · 空文字 → 该块/条目删掉（清空即删除，避免留下空壳）。
    """
    slot = str(slot or "").strip()
    m = _SLOT_RE.match(slot)
    if not m:
        return {"ok": False, "error": f"不认识的槽位：{slot or '(空)'}"}
    text = str(text if text is not None else "").strip()
    if text and len(text) > 400:                     # 与 clean_slides 的上限一致
        text = text[:400]
    base = m.group("base")

    if base in ("title", "subtitle", "meta", "label", "closing", "caption", "index", "cover_title"):
        # 普通页的标题字段是 `name`，封面/结尾页是 `title`；两个都写，渲染时按版式取
        if base == "title":
            old = _text(slide.get("name") or slide.get("title"))
            slide["title"] = text
            slide["name"] = text
            # 封面页的标题可能另存在 `cover_title` 里（渲染时优先取它）——
            # 它原来跟标题是同一串字时**一起改**，否则替换完封面还是老字（踩过）；
            # 老师特意写的不一样的封面标题不动它。
            if _text(slide.get("cover_title")) and _text(slide["cover_title"]) == old:
                slide["cover_title"] = text
        else:
            slide[base] = text
        return {"ok": True, "slot": slot, "text": text}

    if base.startswith("block:"):
        i = int(m.group("bi"))
        blocks = _blocks(slide)
        while len(blocks) <= i:
            blocks.append({"label": "", "text": ""})     # 下标补齐到该位置（占位块）
        if slot.endswith(":label"):
            blocks[i]["label"] = text
        else:
            blocks[i]["text"] = text
            if not text and not _text(blocks[i].get("label")):
                del blocks[i]                            # 清空即删（但**只删被清的那一块**）
        slide["blocks"] = blocks
        slide.pop("bullets", None)                   # 统一走 blocks，避免两处数据打架
        return {"ok": True, "slot": slot, "text": text}

    side = m.group("side")
    col = slide.get(side) if isinstance(slide.get(side), dict) else {}
    heading, items = _text(col.get("heading")), [_text(x) for x in (col.get("bullets") or []) if _text(x)]
    if slot.endswith(":heading"):
        heading = text
    else:
        j = int(m.group("ci"))
        while len(items) <= j:
            items.append("")
        if text:
            items[j] = text
        else:
            del items[j]
    if heading or items:
        slide[side] = {"heading": heading, "bullets": [x for x in items if x]}
    else:
        slide.pop(side, None)
    return {"ok": True, "slot": slot, "text": text}


def apply_override(slide: dict, slot: str, patch: dict) -> dict:
    """把"这一格被拖动/改过大小/改过层级"并进 `slide["slots"]`（R14 的写入口）。

    与 `apply_slot`（改**文字**）分开：文字进 `blocks/bullets/…` 字段，样子进 `slots` 覆盖。
    为什么删字要用 `hidden` 而不是把块删掉：槽位名是**按下标**取的（`block:2`），
    真把第 2 块删了，第 3 块的几何覆盖就会落到新的一块上（老师的拖动成果会"串页"）。
    所以"删除"= `hidden`（不画、字还在、能撤销），索引稳定。

    边界：认不出的槽位/键一律拒绝或丢弃（不猜）；数值由 `_clean_overrides` 夹到范围内；
    `hidden: false` 是**取消隐藏**（把这个键去掉），不是"留着 false"。
    """
    slot = str(slot or "").strip()
    if not _SLOT_RE.match(slot):
        return {"ok": False, "error": f"不认识的槽位：{slot or '(空)'}"}
    if not isinstance(patch, dict) or not patch:
        return {"ok": False, "error": "patch 是空的（要给我 dx/dy/dw/dh/size/color/align/lh/z/hidden）"}
    cur = dict((slide.get("slots") or {}).get(slot) or {})
    cur.update(patch)
    if "hidden" in patch and not patch.get("hidden"):
        cur.pop("hidden", None)
    clean = _clean_overrides({slot: cur}).get(slot) or {}
    slots = dict(slide.get("slots") or {})
    if clean:
        slots[slot] = clean
    else:
        slots.pop(slot, None)                  # 全零覆盖 = 回到版式原位（不留空壳）
    if slots:
        slide["slots"] = slots
    else:
        slide.pop("slots", None)
    return {"ok": True, "slot": slot, "slots": dict(slots)}


def pick_layout(slide: dict, index: int, total: int) -> str:
    """版式：显式 `layout` 优先；否则按位置推断（docs/19 §8.4）。

    ⚠️ 第 1 页**只有在没有要点内容时**才当封面：否则把"情境导入"这类真内容页排成封面，
    它的要点会被整页丢掉（这是"内容零丢失"的红线，宁可不给封面也不丢内容）。
    最后一页用结尾版式（它本来就排要点，不会丢）。
    """
    layout = _text(slide.get("layout"))
    if layout in LAYOUTS:
        return layout
    if index == 0 and not _items(slide):
        return "cover"
    if total > 1 and index == total - 1:
        return "end"
    return "bullets"


def _clean_image(raw) -> dict:
    """收干净"这一页的配图"（R18）：只认 `data/uploads/` 下的**图片相对路径** + 说明与出处。

    为什么这么严：`src` 是**渲染时要去读的文件**（读成 data URI 内联进 SVG）。
    只认"`<会话>/<文件名>.<图片后缀>`"这一种形状 —— 绝对路径、`..`、非图片后缀一律丢；
    真读不到图时 `slide_image.data_uri()` 还会再挡一次（越界路径返回空串）。
    """
    if not isinstance(raw, dict):
        return {}
    src = str(raw.get("src") or "").strip().replace("\\", "/")
    parts = [p for p in src.split("/") if p not in ("", ".")]
    if (not parts or len(parts) > 3 or any(p == ".." for p in parts)
            or len(src) > 200 or Path(parts[-1]).suffix.lower() not in
            (".png", ".jpg", ".jpeg", ".webp", ".gif")):
        return {}
    out = {"src": "/".join(parts)}
    for k, n in (("alt", 80), ("credit", 40)):
        v = _text(raw.get(k))[:n]
        if v:
            out[k] = v
    return out


def label_for(slide: dict, index: int) -> str:
    """内容页左上角的小标签：`环节 2 · 12 分钟`（时长来自教案里的"约 12 分钟"）。"""
    label = _text(slide.get("label"))
    if label:
        return label
    minutes = _text(slide.get("minutes"))
    if not minutes:
        return ""
    return f"环节 {index + 1} · {minutes} 分钟"


def to_slots(slide: dict, index: int, total: int, doc_title: str = "") -> tuple[str, dict]:
    """把一页会话数据映射成 `(版式, 槽位)`。"""
    layout = pick_layout(slide, index, total)
    name = _text(slide.get("name") or slide.get("title"))
    slots: dict = {}

    if layout == "cover":
        # 封面：整份教案的课题当主标，第一个环节名降为副标（比拿环节名当封面好看）
        title = _text(slide.get("cover_title")) or doc_title or name
        subtitle = _text(slide.get("subtitle")) or (name if title != name else "")
        slots.update({"title": title, "subtitle": subtitle,
                      "meta": _text(slide.get("meta")), "label": _text(slide.get("label"))})
    elif layout == "section":
        slots.update({"index": _text(slide.get("index")) or str(index + 1), "title": name})
    elif layout == "two-col":
        left = slide.get("left") if isinstance(slide.get("left"), dict) else {}
        right = slide.get("right") if isinstance(slide.get("right"), dict) else {}
        slots.update({"title": name, "label": label_for(slide, index),
                      "left": {"heading": _text(left.get("heading")),
                               "bullets": [_text(x) for x in (left.get("bullets") or []) if _text(x)]},
                      "right": {"heading": _text(right.get("heading")),
                                "bullets": [_text(x) for x in (right.get("bullets") or []) if _text(x)]}})
    elif layout == "media":
        pic = slide.get("image") if isinstance(slide.get("image"), dict) else {}
        src = _text(pic.get("src"))
        slots.update({"title": name, "label": label_for(slide, index),
                      "caption": _text(slide.get("caption")), "bullets": _items(slide),
                      # 图：**只读** data/uploads 里的文件，现算成 data: URI（预览不落盘）
                      "image": (slide_image.data_uri(src) if src else ""),
                      "image_credit": _text(pic.get("credit"))[:40]})
    elif layout == "end":
        slots.update({"title": name or "小结", "label": label_for(slide, index),
                      "bullets": _items(slide), "closing": _text(slide.get("closing"))})
    else:                                   # bullets
        slots.update({"title": name, "label": label_for(slide, index), "blocks": _blocks(slide)})
    return layout, slots


# ---------------------------------------------------------------- 整册

def canvas_info(t: dict) -> dict:
    """给面板的画布参数（吸附参考线、网格、缩放都用它，**与服务端同一份数值**）。

    `margin/content_w/cols/gutter` 全来自 tokens 的 canvas —— 前端不许自己写死 72 / 12。
    """
    c = t["canvas"]
    return {"w": CANVAS_W, "h": CANVAS_H, "margin": float(c.get("margin") or 72),
            "cols": int(c.get("cols") or 12), "gutter": float(c.get("gutter") or 24),
            "content_w": float((c.get("content_w") or [0, 1136])[1])}


def build_deck(theme_id: str, slides: list[dict] | None, title: str = "",
               tokens: dict | None = None, *, annotate: bool = False) -> dict:
    """渲染整册：返回 `{ok, theme_id, pages:[svg...], layouts, warnings, stats[, boxes]}`。

    **不落盘、不联网、不依赖 WPS**；`pages` 就是面板要显示、导出要转换的那份 SVG。
    `tokens` 给了就直接用（导入模板走这条：借它的配色/字体，版式仍用我们的 6 个）。

    `annotate=True`（**只有预览用**）：给文字槽打 `data-slot` 并额外返回 `boxes`（每页一组可编辑盒子），
    编辑器拿它做页内编辑；导出路径**不传它**，所以导出的 SVG 与以前逐字节一致。
    """
    t = tokens if tokens else theme(theme_id)
    items = [s for s in (slides or []) if isinstance(s, dict)]
    if not items:
        return {"ok": True, "theme_id": t["id"], "pages": [], "layouts": [],
                "boxes": [], "warnings": [], "fonts": font_choices(t),
                "canvas": canvas_info(t),
                "stats": {"pages": 0, "chars": 0}}

    total = len(items)
    warnings: list[str] = list(length_warnings(items))
    pages: list[str] = []
    layouts: list[str] = []
    boxes_all: list[list[dict]] = []
    chars = 0
    for i, slide in enumerate(items):
        layout, slots = to_slots(slide, i, total, doc_title=_text(title))
        layouts.append(layout)
        page_boxes: list[dict] | None = [] if annotate else None
        pages.append(ppt_svg.render_slide(t, layout, slots, index=i, total=total, warnings=warnings,
                                          boxes=page_boxes, overrides=slide.get("slots") or None,
                                          objects=slide.get("objects") or None))   # R21 自由层
        if page_boxes is not None:
            boxes_all.append(page_boxes)
        chars += sum(len(str(v)) for v in (slide.get("bullets") or []))
    out = {"ok": True, "theme_id": t["id"], "pages": pages, "layouts": layouts,
           "warnings": warnings,
           "canvas": canvas_info(t),
           "fonts": font_choices(t),          # R15：右栏字体下拉（主题栈在前 + 常见中文字体）
           "stats": {"pages": total, "chars": chars, "theme_name": t.get("name") or t["id"]}}
    if annotate:
        out["boxes"] = boxes_all
    return out


def slide_warnings(theme_id: str, slides: list[dict] | None, title: str = "") -> list[str]:
    """只跑"密度/上限"检查（面板提示用），不返回 SVG。"""
    return build_deck(theme_id, slides, title).get("warnings") or []


#: 画布尺寸（与 tokens 一致；面板的等比缩放以此为准）
CANVAS_W = 1280
CANVAS_H = 720
#: 缩放系数上下限（太小看不清、太大浪费屏幕）
SCALE_MIN, SCALE_MAX = 0.15, 2.0
#: 每块建议字数（超出只提示，见 `length_warnings`）
CHARS_PER_BLOCK = 22


def fit_scale(container_w: float) -> float:
    """面板宽度 → 预览台缩放系数（**面板里的 JS 是这段公式的镜像**）。

    预览台内部永远是 1280×720；只有这一个**均匀**缩放，所以元素之间的相对位置与导出完全一致
    （面板拖宽拖窄都不会变形）。单测锁住取值，面板检查锁住 `data-scale`。
    """
    try:
        w = float(container_w)
    except (TypeError, ValueError):
        w = 0.0
    if w <= 0:
        w = 480.0                      # 量不到宽度时的兜底（jsdom / 隐藏时就是这种）
    return max(SCALE_MIN, min(SCALE_MAX, w / CANVAS_W))


def fit_height(container_w: float) -> int:
    """缩放后预览台应有的高度（面板据此设 wrap 的高度，避免留白/裁切）。

    用 `floor(x+0.5)` 而不是 `round()`：JS 的 `Math.round(382.5)=383`，Python 的 `round(382.5)=382`，
    两边不一致就会差 1px —— 审计算出来的高度必须与面板实际设的一致。
    """
    return int(math.floor(CANVAS_H * fit_scale(container_w) + 0.5))


#: 老师改过的样子（`slides[i].slots`）里认的键与取值范围（R14/R15）：
#: 位置/大小是"相对版式原位的偏移"（设计像素），字号/颜色/字体/对齐是语义覆盖，hidden = 删掉这一处，
#: `z` = 层级（正数往上盖，画的时候往后排），`bold`/`italic` = 加粗/斜体。
#: **认不出的一律丢**（不悄悄改字、也不让越界值进来 —— 越界由审计去报，不在这里静默夹掉太狠）。
SLOT_OVERRIDE = {
    "dx": (-200.0, 200.0), "dy": (-200.0, 200.0),
    "dw": (-400.0, 400.0), "dh": (-400.0, 400.0),
    "size": (8.0, 200.0), "lh": (0.8, 3.0),
}
#: 层级上下限：老师拖一块到另一块上面，用「上移一层/下移一层」调，值就是整数层号
SLOT_Z_RANGE = (-9, 9)
SLOT_ALIGNS = ("start", "middle", "end")
#: 右栏字体清单（R15）：前 3 个是主题自带的栈，其余是 Windows/常见中文字体。
#: **一份来源**给前后端用（前端从这个列表出下拉，服务端只认这个列表里的 + 主题栈）。
SLOT_FONTS = ("微软雅黑", "黑体", "宋体", "楷体", "等线", "思源黑体", "Noto Sans SC")
#: 字体名上限（防手滑粘一整段）
FONT_MAX_CHARS = 40


def clean_font(name: object) -> str:
    """字体名收口：去空白、限长；空/超长一律当"没设"（服务端不猜）。"""
    s = str(name or "").strip().strip('"\'')
    if not s or len(s) > FONT_MAX_CHARS:
        return ""
    return s if all(ch not in s for ch in ("<", ">", "&", ";")) else ""


def font_choices(theme: dict | None = None) -> list[str]:
    """右栏字体下拉的清单：**主题栈在前**（默认值来源），其余是常见中文字体（去重）。"""
    out: list[str] = []
    for f in ((theme or {}).get("fonts") or {}).get("stack") or []:
        f = clean_font(f)
        if f and f.lower() not in ("sans-serif", "serif", "arial") and f not in out:
            out.append(f)
    for f in SLOT_FONTS:
        if f not in out:
            out.append(f)
    return out


def _clean_overrides(raw) -> dict:
    """收干净 `slots` 覆盖：只认已知槽位名与已知键，数值夹进 `SLOT_OVERRIDE` 的范围。"""
    if not isinstance(raw, dict):
        return {}
    out: dict = {}
    for slot, val in raw.items():
        slot = str(slot or "").strip()
        if not slot or not _SLOT_RE.match(slot) or not isinstance(val, dict):
            continue
        item: dict = {}
        for k, (lo, hi) in SLOT_OVERRIDE.items():
            if k in val:
                try:
                    raw = float(val[k])
                except (TypeError, ValueError):
                    continue
                if not raw:
                    continue           # 0 = "没改"（**不是**"夹到下限"：size:0 曾被夹成 8px —— 踩过）
                item[k] = round(max(lo, min(hi, raw)), 1)
        if "z" in val:
            try:
                z = int(round(float(val["z"])))
            except (TypeError, ValueError):
                z = 0
            if z:
                item["z"] = max(SLOT_Z_RANGE[0], min(SLOT_Z_RANGE[1], z))
        if val.get("align") in SLOT_ALIGNS:
            item["align"] = val["align"]
        m = re.fullmatch(r"#?([0-9a-fA-F]{6})", str(val.get("color") or ""))
        if m:
            item["color"] = "#" + m.group(1).upper()
        font = clean_font(val.get("font"))
        if font:
            item["font"] = font
        for flag in ("bold", "italic", "underline"):   # R15/R18：加粗/斜体/下划线（Ctrl+B/I/U）
            if val.get(flag):
                item[flag] = True
        if val.get("hidden"):
            item["hidden"] = True
        for k in [k for k, v in item.items() if k in SLOT_OVERRIDE and not float(v or 0)]:
            item.pop(k)                    # 偏移/层级为 0 = 没改，别在会话库里攒一堆 0
        if item:
            out[slot] = item
    return out


# ---------------------------------------------------------------- 自由层对象（R21）

#: 一页最多几个自由层对象（防手滑粘一堆；超出的丢掉并如实说）
MAX_OBJECTS = 30
#: 支持的形状（与 `ppt_svg.SHAPE_PRESETS` 一致；都是 ppt-master 认的 DrawingML 预设）
OBJECT_SHAPES = ppt_svg.SHAPE_PRESETS
#: 几何范围（设计像素 1280×720；允许拖到画布外一点，但不能跑到天边）
OBJ_RANGE = {"x": (-2000.0, 3000.0), "y": (-2000.0, 3000.0), "w": (16.0, 2400.0), "h": (16.0, 2400.0)}
OBJ_ID_RE = re.compile(r"^[A-Za-z0-9_-]{1,24}$")


def _hex_color(v) -> str:
    """认 `#RGB` / `#RRGGBB`（大小写都行），统一成 `#RRGGBB` 大写；认不出返回空串（不猜）。"""
    s = str(v or "").strip()
    m = re.fullmatch(r"#?([0-9a-fA-F]{6})", s)
    if m:
        return "#" + m.group(1).upper()
    m = re.fullmatch(r"#?([0-9a-fA-F]{3})", s)
    if m:
        return "#" + "".join(ch * 2 for ch in m.group(1)).upper()
    return ""


def clean_objects(raw) -> list[dict]:
    """收干净一页的自由层对象（R21）：只认认识的形状/键，数值夹范围，**不认识的直接丢**。

    为什么要单独收一遍：这些是老师**自己造**的东西（不是版式给的槽位），一旦写进会话就会
    长期存在 —— 渲染、导出、审计三处都吃它，所以入口必须严。
    """
    out: list[dict] = []
    seen: set[str] = set()
    for i, o in enumerate(raw or []):
        if not isinstance(o, dict):
            continue
        if len(out) >= MAX_OBJECTS:
            break
        kind = str(o.get("kind") or "shape")
        if kind not in ("shape", "text"):
            continue
        item: dict = {"kind": kind}
        oid = str(o.get("id") or "").strip()
        if not OBJ_ID_RE.match(oid) or oid in seen:
            n = i + 1                       # 缺 id / id 非法 / id 撞车：都按位置重编一个
            while f"o{n}" in seen:
                n += 1
            oid = f"o{n}"
        seen.add(oid)
        item["id"] = oid
        if kind == "shape":
            shape = str(o.get("shape") or "rect")
            item["shape"] = shape if shape in OBJECT_SHAPES else "rect"
        for k, (lo, hi) in OBJ_RANGE.items():
            try:
                v = float(o.get(k))
            except (TypeError, ValueError):
                v = 0.0
            item[k] = round(max(lo, min(hi, v)), 1)
        if item["w"] < OBJ_RANGE["w"][0]:
            item["w"] = 120.0
        if item["h"] < OBJ_RANGE["h"][0]:
            item["h"] = 60.0
        text = _text(o.get("text"))[:400]
        if text:
            item["text"] = text
        try:
            size = float(o.get("size"))
        except (TypeError, ValueError):
            size = 0.0
        if size >= 8:
            item["size"] = round(max(8.0, min(200.0, size)), 1)
        for key in ("color", "fill", "stroke"):
            c = _hex_color(o.get(key))
            if c:
                item[key] = c
        if o.get("align") in SLOT_ALIGNS:
            item["align"] = o["align"]
        font = clean_font(o.get("font"))
        if font:
            item["font"] = font
        for flag in ("bold", "italic", "underline"):
            if o.get(flag):
                item[flag] = True
        try:
            lh = float(o.get("lh"))
        except (TypeError, ValueError):
            lh = 0.0
        if lh:
            item["lh"] = round(max(0.8, min(3.0, lh)), 2)
        try:
            r = float(o.get("radius"))
        except (TypeError, ValueError):
            r = 0.0
        if r:
            item["radius"] = round(max(0.0, min(200.0, r)), 1)
        out.append(item)
    return out


def new_object(kind: str = "shape", shape: str = "rect", *, canvas: tuple[float, float] = (1280, 720),
               theme: dict | None = None, oid: str = "") -> dict:
    """新插入一个对象（默认摆中间、给个能看的默认配色）—— 前端点一下就调这里。

    默认外观守"瑞士风"硬规则（`docs/26` §11.2）：**直角、纯色、无渐变无阴影**；
    颜色取自主题（`soft` 底 + `accent` 描边），不自己编颜色。
    """
    W, H = float(canvas[0] or 1280), float(canvas[1] or 720)
    pal = ((theme or {}).get("palette") or {})
    soft = _hex_color(pal.get("soft")) or "#EEF2F8"
    accent = _hex_color(pal.get("accent")) or "#2F6BFF"
    if kind == "text":
        w, h = 420.0, 90.0
        obj = {"kind": "text", "x": round((W - w) / 2, 1), "y": round((H - h) / 2, 1), "w": w, "h": h,
               "text": "点这里改字", "size": 28.0, "color": _hex_color(pal.get("text")) or "#1A1A1A",
               "align": "start"}
        if oid:
            obj["id"] = oid
        return clean_objects([obj])[0]
    w, h = 320.0, 160.0
    obj = {"kind": "shape", "shape": shape if shape in OBJECT_SHAPES else "rect",
           "x": round((W - w) / 2, 1), "y": round((H - h) / 2, 1), "w": w, "h": h,
           "fill": soft, "stroke": accent, "stroke_w": 2.0}
    if oid:
        obj["id"] = oid
    return clean_objects([obj])[0]


def next_object_id(objects: list[dict] | None) -> str:
    """给这一页的下一个新对象取个不撞车的 id（`o1`、`o2`…）。"""
    used = {str(o.get("id") or "") for o in (objects or []) if isinstance(o, dict)}
    n = 1
    while f"o{n}" in used:
        n += 1
    return f"o{n}"


def patch_object(objects: list[dict] | None, oid: str, patch: dict | None) -> dict:
    """改一个自由层对象（几何/文字/样式），返回 `{ok, objects, obj}`。

    **几何是绝对设计像素**（`x/y/w/h` 直接给），不是槽位那套 `dx/dy` 增量 ——
    自由层对象本来就没有"版式原位"可言，增量语义只会让"拖回原地"变得说不清。
    """
    if not isinstance(patch, dict):
        return {"ok": False, "error": "补丁不是个对象"}
    items = [dict(o) for o in (objects or []) if isinstance(o, dict)]
    hit = None
    for o in items:
        if str(o.get("id")) == str(oid):
            hit = o
            break
    if hit is None:
        return {"ok": False, "error": f"这一页没有对象 {oid}"}
    for k in ("x", "y", "w", "h", "size", "lh", "radius", "stroke_w"):
        if k in patch and patch[k] is not None:
            try:
                hit[k] = float(patch[k])
            except (TypeError, ValueError):
                return {"ok": False, "error": f"{k} 得是数字"}
    if "text" in patch:
        hit["text"] = str(patch["text"] or "")[:400]
    for k in ("color", "fill", "stroke"):
        if k in patch:
            c = _hex_color(patch[k])
            if c:
                hit[k] = c
            elif str(patch[k] or "") in ("", "none"):
                hit.pop(k, None)
    if "align" in patch and patch["align"] in SLOT_ALIGNS:
        hit["align"] = patch["align"]
    if "font" in patch:
        f = clean_font(patch["font"])
        if f:
            hit["font"] = f
        else:
            hit.pop("font", None)
    for flag in ("bold", "italic", "underline"):
        if flag in patch:
            if patch[flag]:
                hit[flag] = True
            else:
                hit.pop(flag, None)
    if "shape" in patch and str(patch["shape"]) in OBJECT_SHAPES:
        hit["shape"] = str(patch["shape"])
    cleaned = clean_objects(items)          # 夹范围、去非法值；顺便保证 id 不撞车
    got = next((o for o in cleaned if str(o.get("id")) == str(oid)), None)
    if got is None:
        return {"ok": False, "error": "改完这个对象就不合法了（检查一下尺寸/颜色）"}
    return {"ok": True, "objects": cleaned, "obj": got}


def clean_slides(slides: list[dict] | None) -> list[dict]:
    """把面板传来的编辑结果**收干净**再存/再渲染（保守：不认识的就丢，绝不悄悄改字）。

    只保留认识的键：`layout/name/title/minutes/label/bullets/blocks/caption/closing/
    subtitle/meta/index/left/right/objects/image/notes/slots`；字符串去首尾空白；
    `layout` 只认登记过的 6 个；空条目丢掉；单块文字上限 400 字（防手滑粘贴一整篇）。
    """
    out: list[dict] = []
    for raw in slides or []:
        if not isinstance(raw, dict):
            continue
        s: dict = {}
        for key in ("name", "title", "minutes", "label", "caption", "closing",
                    "subtitle", "meta", "index", "cover_title"):
            if key in raw:
                s[key] = _text(raw[key])[:120]
        layout = _text(raw.get("layout"))
        if layout in LAYOUTS:
            s["layout"] = layout
        bullets = [_text(b)[:400] for b in (raw.get("bullets") or []) if _text(b)]
        if bullets:
            s["bullets"] = bullets
        blocks = []
        for b in (raw.get("blocks") or []):
            if isinstance(b, dict):
                text = _text(b.get("text") or b.get("name"))[:400]
                if text:
                    blocks.append({"label": _text(b.get("label"))[:24], "text": text})
            elif _text(b):
                blocks.append({"label": "", "text": _text(b)[:400]})
        if blocks:
            s["blocks"] = blocks
        for side in ("left", "right"):
            col = raw.get(side)
            if isinstance(col, dict):
                items = [_text(x)[:400] for x in (col.get("bullets") or []) if _text(x)]
                if items or _text(col.get("heading")):
                    s[side] = {"heading": _text(col.get("heading"))[:60], "bullets": items}
        slots = _clean_overrides(raw.get("slots"))          # R14：老师拖过/改过的样子
        if slots:
            s["slots"] = slots
        pic = _clean_image(raw.get("image"))                # R18：这一页的配图（联网取回来的）
        if pic:
            s["image"] = pic
        objs = clean_objects(raw.get("objects"))            # R21：老师自己插的文本框/形状
        if objs:
            s["objects"] = objs
        notes = _text(raw.get("notes"))[:2000]              # R16 预留：每一页的备注
        if notes:
            s["notes"] = notes
        if s.get("name") or s.get("title") or s.get("bullets") or s.get("blocks") or s.get("left"):
            out.append(s)                       # 全空的一页直接丢，别留下空白页
    return out


def length_warnings(slides: list[dict] | None) -> list[str]:
    """逐块字数提示（T4.2）：超上限只提示，**不改写、不截断**（与 §3.3 一致）。"""
    out: list[str] = []
    for i, s in enumerate(slides or []):
        blocks = _blocks(s) or [{"label": "", "text": b} for b in _bullets(s)]
        for j, b in enumerate(blocks, 1):
            if len(b["text"]) > CHARS_PER_BLOCK:
                out.append(f"第 {i + 1} 页第 {j} 块 {len(b['text'])} 字 > {CHARS_PER_BLOCK} 字"
                           f"（「{b['text'][:12]}…」）→ 建议精简或拆页")
    return out


# ---------------------------------------------------------------- 导出（docs/23 M3）

#: 导出后端开关：不设/`native` = 走 ppt-master（新，原生可编辑）；`legacy` = 回退旧版式导出
ENV_EXPORT = "EDU_PPT_EXPORT"
#: 允许作为中文字体的族名（本机已装：微软雅黑 / Noto Sans SC）
EA_OK = ("微软雅黑", "Microsoft YaHei", "Noto Sans SC", "Source Han Sans", "思源黑体")


def export_backend() -> str:
    """当前导出后端：`native`（默认）或 `legacy`（一键回退，见 docs/22 §6.3）。"""
    v = (os.environ.get(ENV_EXPORT) or "").strip().lower()
    return "legacy" if v in ("legacy", "old", "0", "off") else "native"


def safe_name(s: str, fallback: str = "教案") -> str:
    """文件名安全化（老师给的课题里什么都可能有）。"""
    bad = '<>:"/\\|?*'
    out = "".join("_" if ch in bad else ch for ch in str(s or "")).strip().strip(".")
    return (out or fallback)[:40]


def svg_texts(svg: str) -> list[str]:
    """抠出一页 SVG 里的可见文字（导出后比对内容没丢）。"""
    return [html.unescape(t) for t in re.findall(r"<text[^>]*>(.*?)</text>", svg, re.S) if t.strip()]


def _has_cjk(s: str) -> bool:
    return any(ord(ch) > 0x2E80 for ch in s)


def pptx_font_audit(slides_xml: list[str], theme_xml: str = "") -> tuple[bool, str]:
    """查"中文字体有没有掉回宋体"：**只看含中文的 run**（纯西文 run 的 ea 无害）。

    转换器会给纯西文 run 写 `<a:ea typeface="Arial">`，这不是问题；
    真会出事的是含中文的 run 没写 ea —— 那时渲染器按主题缺省走，中文 Windows 上就是宋体。
    """
    total = explicit = 0
    bad: list[str] = []
    eas: set[str] = set()
    latins: set[str] = set()
    for xml in slides_xml:
        for run in re.findall(r"<a:r>(.*?)</a:r>", xml, re.S):
            txt = html.unescape("".join(re.findall(r"<a:t>([^<]*)</a:t>", run)))
            lat = re.search(r'<a:latin typeface="([^"]*)"', run)
            ea = re.search(r'<a:ea typeface="([^"]*)"', run)
            if lat and lat.group(1):
                latins.add(lat.group(1))
            if not _has_cjk(txt):
                continue
            total += 1
            if ea and ea.group(1):
                explicit += 1
                eas.add(ea.group(1))
                if ea.group(1) not in EA_OK:
                    bad.append(f"「{txt[:8]}」→ ea={ea.group(1)}")
            else:
                bad.append(f"「{txt[:8]}」→ 没写 ea（会跟主题缺省，中文环境多半掉宋体）")
    theme_ea = ""
    if theme_xml:
        m = re.search(r"<a:minorFont>.*?<a:ea typeface=\"([^\"]*)\"", theme_xml, re.S)
        theme_ea = (m.group(1) if m else "") or "（空 = 跟随系统默认）"
    ok = total > 0 and not bad
    msg = (f"含中文 run {total} 个，显式 ea {explicit} 个 → {sorted(eas) or '—'}"
           f"；主题缺省 ea={theme_ea}；西文 latin={sorted(latins) or '—'}")
    if bad:
        msg += "；问题：" + "；".join(bad[:3])
    return ok, msg


def slide_notes(slides: list[dict] | None, count: int = 0) -> list[str]:
    """每页的讲稿备注（R16）：与页一一对应，缺的补空串（导出时按页写进 notesSlide）。"""
    out = []
    for s in (slides or []):
        out.append(_text((s or {}).get("notes"))[:2000] if isinstance(s, dict) else "")
    while len(out) < count:
        out.append("")
    return out


def notes_in_pptx(pptx: Path) -> list[str]:
    """读回 pptx 里每页的备注（校验用；读不到就是空串）。

    为什么不直接读包里的 XML：python-pptx 已经把 notesSlide 的坑（缺少 notesMaster、
    空段落、占位符继承）都处理过了，读它的 `notes_text_frame.text` 最省事也最准。
    """
    try:
        from pptx import Presentation
        prs = Presentation(str(pptx))
    except Exception:  # noqa: BLE001 读不出来就当没有备注（校验会如实报"缺"）
        return []
    out: list[str] = []
    for slide in prs.slides:
        try:
            out.append((slide.notes_slide.notes_text_frame.text or "").strip())
        except Exception:  # noqa: BLE001 没有备注页 / 模板异常
            out.append("")
    return out


def write_notes(pptx: Path, notes: list[str]) -> dict:
    """把每页备注写进 pptx（R16）：**就地改产物**，写不进去如实返回，不假装成功。

    为什么放在转换之后：ppt-master 是 svg→pptx，SVG 里没有"备注"这个概念，
    只能在产物上补 —— 补完还要 `notes_in_pptx` 读回来核对（校验里那一条就是它）。
    """
    from pptx import Presentation

    want = [(n or "").strip() for n in notes]
    if not any(want):
        return {"ok": True, "written": 0, "detail": "没有备注要写"}
    try:
        prs = Presentation(str(pptx))
    except Exception as exc:  # noqa: BLE001
        return {"ok": False, "written": 0, "error": f"{type(exc).__name__}: {exc}"}
    if len(prs.slides) != len(want):
        return {"ok": False, "written": 0,
                "error": f"页数不一致：pptx {len(prs.slides)} 页 / 备注 {len(want)} 条"}
    written = 0
    for slide, text in zip(prs.slides, want):
        if not text:
            continue
        try:
            slide.notes_slide.notes_text_frame.text = text
            written += 1
        except Exception as exc:  # noqa: BLE001 单页失败不拖累别的页
            return {"ok": False, "written": written, "error": f"{type(exc).__name__}: {exc}"}
    prs.save(str(pptx))
    return {"ok": True, "written": written, "detail": f"写了 {written} 页备注"}


def verify_pptx(pptx: Path, pages: list[str], notes: list[str] | None = None) -> dict:
    """`docs/22` §6.2 的四条硬验收：页数 / 原生文本 / 不嵌字体+中文字体 / 内容零丢失。

    返回 `{ok, checks:[{name, ok, detail}], detail}`；**只看报告与包内容，不看退出码**
    （`svg_to_pptx.py` 在 `passed-with-warnings` 时退出码仍是 1）。

    `notes` 给了就多一条 **备注读回**（R16）：写了讲稿就要能在 pptx 里读出来，
    读不回来说明白写了（不能只看"我们调用赋值没报错"）。
    """
    checks: list[dict] = []

    def add(name: str, ok: bool, detail: str = "") -> None:
        checks.append({"name": name, "ok": bool(ok), "detail": detail})

    if not pptx.is_file():
        add("产物存在", False, str(pptx))
        return {"ok": False, "checks": checks, "detail": "没有产出 pptx"}

    z = zipfile.ZipFile(pptx)
    parts = z.namelist()
    slides = sorted(n for n in parts if re.fullmatch(r"ppt/slides/slide\d+\.xml", n))
    per_slide, shapes, slides_xml = [], 0, []
    for name in slides:
        xml = z.read(name).decode("utf-8")
        slides_xml.append(xml)
        per_slide.append([html.unescape(t) for t in re.findall(r"<a:t>([^<]*)</a:t>", xml)])
        shapes += len(re.findall(r"<p:sp>", xml))
    theme_xml = ""
    for name in parts:
        if re.fullmatch(r"ppt/theme/theme\d+\.xml", name):
            theme_xml = z.read(name).decode("utf-8")
            break
    embedded = [n for n in parts if "fntdata" in n]
    want = [svg_texts(p) for p in pages]

    font_ok, font_msg = pptx_font_audit(slides_xml, theme_xml)
    add("页数一致", len(slides) == len(pages), f"pptx {len(slides)} 页 / SVG {len(pages)} 页")
    add("原生文本（不是图片）", bool(shapes) and all(per_slide),
        f"{shapes} 个形状，每页文字块 {[len(p) for p in per_slide]}")
    add("不嵌字体（WPS 不读内嵌裸 TTF）", not embedded, f"包内 fntdata {len(embedded)} 个")
    add("中文字体显式指定", font_ok, font_msg)
    missing = []
    for i, (w, got) in enumerate(zip(want, per_slide), 1):
        lack = [t for t in w if t.strip() and t.strip() not in [g.strip() for g in got]]
        if lack:
            missing.append(f"第 {i} 页缺 {lack[:3]}")
    add("内容零丢失", not missing and len(want) == len(per_slide),
        "逐页逐串比对通过" if not missing else "；".join(missing))

    if notes is not None:
        want_notes = [(n or "").strip() for n in notes]
        got_notes = notes_in_pptx(pptx)
        lack = [i + 1 for i, (w, g) in enumerate(zip(want_notes, got_notes))
                if w and w not in g]
        short = len(got_notes) < len(want_notes)
        add("备注读回", not lack and not short,
            f"{sum(1 for w in want_notes if w)} 页有备注，逐页读回通过" if not lack and not short
            else (f"备注页数不够：{len(got_notes)} / {len(want_notes)}" if short
                  else f"这些页的备注读不回来：{lack[:5]}"))

    ok = all(c["ok"] for c in checks)
    return {"ok": ok, "checks": checks,
            "detail": "；".join(f"{'✅' if c['ok'] else '❌'} {c['name']}" for c in checks)}


def export_deck_pptx(theme_id: str, slides: list[dict] | None, title: str = "", *,
                     out_dir: str | Path | None = None) -> dict:
    """把幻灯片导出成**原生可编辑 pptx**（同一份 SVG 交给 ppt-master 转换）。

    落盘规则（`docs/23` §5）：中间件进 `data/tmp/ppt_jobs/<job>/`（24h 自清），
    产物进 `data/ppt_out/`（或老师选的目录）；**`content/plans/` 一个字不写**。
    """
    from . import ppt_polish                       # 复用它的 home/scripts/workroot/outdir

    built = build_deck(theme_id, slides, title=title)
    pages = built["pages"]
    if not pages:
        return {"ok": False, "error": "没有可导出的页（slides 为空？）", "backend": "native"}

    scripts = ppt_polish.scripts_dir()
    if not (scripts / "svg_to_pptx.py").is_file():
        return {"ok": False, "backend": "native", "gate": "install",
                "error": f"ppt-master 未就绪：{scripts}（跑 deploy\\setup_ppt_master.cmd）"}

    try:
        ppt_polish.cleanup_stale()
    except Exception:  # noqa: BLE001 清理失败不该挡住导出
        pass

    stamp = time.strftime("%Y%m%d-%H%M%S")
    job = ppt_polish.workroot() / f"deck-{stamp}"
    svg_out = job / "svg_output"
    svg_out.mkdir(parents=True, exist_ok=True)
    for i, svg in enumerate(pages, 1):
        (svg_out / f"{i:02d}_{built['layouts'][i - 1]}.svg").write_text(svg, encoding="utf-8")

    # ⚠️ 必须 `resolve()`：子进程的 cwd 是任务目录，相对路径会落到 `任务目录/data/tmp/...` 里（踩过）
    target_dir = Path(out_dir).expanduser().resolve() if out_dir else ppt_polish.outdir()
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / f"教案_{safe_name(title or '未命名')}_{stamp}.pptx"
    if target.exists():                      # 同一秒导两次别互相覆盖
        for n in range(2, 100):
            alt = target.with_name(f"{target.stem}-{n}{target.suffix}")
            if not alt.exists():
                target = alt
                break

    logs: list[str] = []
    steps = (
        ([sys.executable, str(scripts / "svg_quality_checker.py"), str(job),
          "--canonical-authoring", "--stage", "final", "--quick-generate", "--json"], "质量检查"),
        ([sys.executable, str(scripts / "svg_to_pptx.py"), str(job),
          "--quick-generate", "-o", str(target)], "转换"),
    )
    for cmd, tag in steps:
        try:
            proc = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8",
                                  errors="replace", timeout=600, cwd=str(job))
        except (OSError, subprocess.TimeoutExpired) as exc:
            return {"ok": False, "backend": "native", "job": str(job),
                    "error": f"{tag}没跑起来：{exc}"}
        tail = [ln for ln in (proc.stdout or "").splitlines() if ln.strip()]
        logs.append(f"{tag}：{(tail[-1] if tail else '')[:120]}")
        if tag == "转换" and not target.is_file():
            return {"ok": False, "backend": "native", "job": str(job), "logs": logs,
                    "error": f"转换未产出文件（exit={proc.returncode}）"}

    report = job / "validation" / f"{target.stem}.report.json"
    status = ""
    if report.is_file():
        try:
            status = json.loads(report.read_text(encoding="utf-8")).get("status", "")
        except Exception:  # noqa: BLE001 报告坏了不当失败
            status = ""
    # R16：备注在转换之后补进产物（SVG 里没有"备注"这个概念），补完由 verify 读回核对
    notes = slide_notes(slides, len(pages))
    post = _postcheck(target, pages, notes)             # 内核：清动画 + 补 <a:ea> + 内容零丢失（必经）
    verify = post.pop("verify", {})
    ok = verify.get("ok", False) and (not status or status.startswith("passed"))
    return {"ok": ok, "backend": "native", "file": str(target), "filename": target.name,
            "outDir": str(target_dir), "job": str(job), "status": status,
            "pages": len(pages), "theme_id": built["theme_id"], "layouts": built["layouts"],
            "verify": verify, "postcheck": post, "logs": logs,
            "notes_written": sum(1 for n in notes if n),
            "error": "" if ok else (f"转换报告 status={status or '（缺失）'}；" + verify.get("detail", ""))}


def _postcheck(path: Path, pages: list[str], notes: list[str] | None = None) -> dict:
    """生成链路的**最后一道**（docs/23 §6.5 R6-4）：内核 + 内容零丢失。失败不影响产物交付，但要如实报出。

    `notes`（R16）先写进产物再校验 —— 顺序反了会校验出一条假失败。
    """
    from . import ppt_polish

    out: dict = {}
    try:
        out["normalize"] = ppt_polish.normalize_file(path)
    except Exception as exc:  # noqa: BLE001 内核失败不该让产物不可用
        out["normalize"] = {"ok": False, "error": f"{type(exc).__name__}: {exc}"}
    if notes is not None:
        try:
            out["notes"] = write_notes(path, notes)
        except Exception as exc:  # noqa: BLE001 备注写不进去不该让产物不可用（但要如实报）
            out["notes"] = {"ok": False, "error": f"{type(exc).__name__}: {exc}"}
    out["verify"] = verify_pptx(path, pages, notes=notes)
    return out


if __name__ == "__main__":        # 冒烟：python -m edu_agent.slide_deck
    import sys

    sys.stdout.reconfigure(encoding="utf-8")
    demo = [{"name": "3.3 幂函数"}, {"name": "情境导入", "minutes": "5", "bullets": ["投影五个实例"]},
            {"name": "小结", "minutes": "3", "bullets": ["三条结论"]}]
    for tid in tk.available():
        out = build_deck(tid, demo, title="3.3 幂函数")
        print(f"{tid}: {out['stats']['pages']} 页 · 版式 {out['layouts']} · 警告 {len(out['warnings'])}")
